"""Scan recalced workbook copy for formula errors; verify new tabs."""
from openpyxl import load_workbook

path = r"C:\Users\Ben\Documents\GitHub\Bipedal_Robot\tmp\recalc_check\Results_table_20mm.xlsx"
wbv = load_workbook(path, data_only=True)
wbf = load_workbook(path, data_only=False)

ERRS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!")
n_err = 0
for name in wbv.sheetnames:
    wsv, wsf = wbv[name], wbf[name]
    for row in wsv.iter_rows():
        for c in row:
            if isinstance(c.value, str) and any(e in c.value for e in ERRS):
                n_err += 1
                print(f"ERROR {name}!{c.coordinate}: {c.value}  formula={wsf[c.coordinate].value!r}")
print("total formula errors:", n_err)

for name in ["ExtTest20mm_1", "FlxTest20mm_51cm"]:
    wsv = wbv[name]
    vals = [wsv.cell(row=16, column=c).value for c in range(3, 25)]
    formulas_ok = all(
        wbf[name].cell(row=16, column=c).value is not None for c in range(3, 25)
    )
    print(f"{name}: row16 cached all-zero={all(v == 0 for v in vals)}, "
          f"values sample={vals[:4]}, formulas present={formulas_ok}, "
          f"charts={len(wbf[name]._charts)}")

# sanity: template tab formulas still evaluate to the same cached values as before
tpl = wbv["FlxTest20mm_42cm (3)"]
print("template C16 still:", tpl["C16"].value, "(expect -5.626...)")
