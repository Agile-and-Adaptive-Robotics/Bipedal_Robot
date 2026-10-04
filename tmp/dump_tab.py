"""Dump every cell (value/formula + key style) of a sheet in Results_table_20mm.xlsx."""
import sys
from openpyxl import load_workbook

path = r"C:\Users\Ben\Documents\GitHub\Bipedal_Robot\Testing_Data\2026_06_Festo\Results_table_20mm.xlsx"
sheet = sys.argv[1] if len(sys.argv) > 1 else "FlxTest20mm_42cm (3)"

wb = load_workbook(path, data_only=False)
ws = wb[sheet]
print(f"=== sheet '{sheet}' dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
print(f"merged: {ws.merged_cells.ranges}")
print("col widths:", {k: round(v.width, 1) if v.width else None for k, v in ws.column_dimensions.items()})

for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
    for c in row:
        if c.value is None:
            continue
        f = c.font
        style = f"{f.name}/{f.size}/{'B' if f.bold else ''}"
        fill = c.fill.start_color.rgb if c.fill and c.fill.fill_type == "solid" else ""
        nf = c.number_format if c.number_format != "General" else ""
        print(f"{c.coordinate:>5} | {repr(c.value)[:80]:<82} | {style:<18} | fill={fill} | nf={nf}")

# also cached values for formula cells
wbv = load_workbook(path, data_only=True)
wsv = wbv[sheet]
print("\n--- cached values of formula cells ---")
for row in ws.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            v = wsv[c.coordinate].value
            print(f"{c.coordinate:>5} | {c.value:<60} -> {v!r}")
