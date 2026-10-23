from pathlib import Path
import openpyxl,json
R=Path(r'D:\Github\Bipedal_Robot');C=R/'Documentation/Reports and Papers/Dissertation/Notes/advisor_revision_20261022/Revision_C'
p=R/'Testing_Data/2026_06_Festo/Results_table_20mm.xlsx'
w=openpyxl.load_workbook(p,data_only=True,read_only=True)
print(w.sheetnames)
records=[]
for s in w:
    if 'Ext' not in s.title:continue
    print('SHEET',s.title,s.max_row,s.max_column)
    if s.title!='ExtTest20mm_2':continue
    for row in s:
        cells={c.coordinate:c.value for c in row if c.value is not None}
        if cells:records.append(cells);print(json.dumps(cells,ensure_ascii=True,default=str))
(C/'review/torque_sheet_evidence.json').write_text(json.dumps({'source':str(p),'sheet':'ExtTest20mm_2','rows':records},indent=2,default=str))
