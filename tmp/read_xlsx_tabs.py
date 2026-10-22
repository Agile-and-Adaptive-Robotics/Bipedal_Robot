import sys

import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
print("TABS:", wb.sheetnames)
for name in sys.argv[2:]:
    if name not in wb.sheetnames:
        print(f"--- {name}: NOT FOUND")
        continue
    ws = wb[name]
    print(f"--- {name} ({ws.max_row}x{ws.max_column}) ---")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40), values_only=True):
        if all(v is None for v in row):
            continue
        cells = ["" if v is None else (f"{v:.4g}" if isinstance(v, float) else str(v)) for v in row]
        print(" | ".join(cells))
