# Independent verifier for the 27-actuator capability-table deliverables.
# Re-derives every number from the CSVs of record; compares md/tex/xlsx cells.
import csv, json, re, sys, statistics
from pathlib import Path
import openpyxl

FOLDER = Path(r"D:\GitHub\Bipedal_Robot\tmp\paper_bridge\capability_table_20261022")
PER = FOLDER / "per_muscle_capability_20261022.csv"
REG = FOLDER / "BiPulley_survey_perMuscle_regraded_20261022_223448.csv"
HTM = FOLDER / "hand_tuned_metrics_20261022.csv"
HTC = FOLDER / "human_torque_curves_20261022.csv"
MD = FOLDER / "capability_table_20261022.md"
TEX = FOLDER / "capability_table_20261022.tex"
XLSX = FOLDER / "capability_table_20261022.xlsx"
SUMJ = FOLDER / "assemble_summary_20261022.json"

problems = []
def check(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg)
    if not cond:
        problems.append(msg)

def num(s):
    s = s.strip()
    return None if s == "" else float(s)

# ---------- 1. per-muscle CSV ----------
with open(PER, newline='', encoding='utf-8') as f:
    per = list(csv.DictReader(f))
print(f"\nper-muscle rows: {len(per)}; ids: {[int(r['actuator_id']) for r in per]}")
check(len(per) == 27, f"per-muscle CSV has 27 rows (got {len(per)})")
check([int(r['actuator_id']) for r in per] == list(range(1, 28)), "actuator_ids are 1..27 in order")

hand = [(num(r['hand_meet_pct']), num(r['hand_active_pct'])) for r in per]
check(all(m is not None and a is not None for m, a in hand), "all 27 hand_meet/hand_active present")
hand_meet_avg = sum(m for m, _ in hand) / len(hand)
hand_act_avg = sum(a for _, a in hand) / len(hand)
print(f"hand: n={len(hand)} meet_avg={hand_meet_avg!r} ({hand_meet_avg:.6f}) act_avg={hand_act_avg!r} ({hand_act_avg:.6f})")

graded = [r for r in per if r['opt_meet_pct'].strip() != ""]
opt_meet_avg = sum(num(r['opt_meet_pct']) for r in graded) / len(graded)
opt_act_avg = sum(num(r['opt_active_pct']) for r in graded) / len(graded)
print(f"opt: n={len(graded)} meet_avg={opt_meet_avg!r} ({opt_meet_avg:.6f}) act_avg={opt_act_avg!r} ({opt_act_avg:.6f})")
check(all(r['opt_active_pct'].strip() != "" for r in graded), "every graded row has both opt cells")

cannot_build = [r for r in per if "CANNOT BUILD" in r['notes']]
grade_unable = [r for r in per if "GRADE-UNABLE" in r['notes']]
built = len(per) - len(cannot_build)
print(f"cannot_build={len(cannot_build)} {[r['muscle'] for r in cannot_build]}")
print(f"grade_unable={len(grade_unable)} {[r['muscle'] for r in grade_unable]}")
print(f"built={built}; graded={len(graded)}; graded+unable={len(graded)+len(grade_unable)}")
check(len(graded) + len(grade_unable) == built, "graded + grade-unable == built count")

suspects = [r for r in per if r['transcription_suspect_flag'].strip().lower() == 'yes']
artifacts = [r for r in per if r['path_artifact_flag'].strip().lower() == 'yes']
print(f"suspects={len(suspects)} {[r['muscle'] for r in suspects]}")
print(f"artifact-flagged={len(artifacts)} {[r['muscle'] for r in artifacts]}")

capped = [r for r in graded if 'capability-capped' in r['notes']]
print(f"capability-capped among graded={len(capped)} {[r['muscle'] for r in capped]}")

unconv_graded = [r for r in graded if r['opt_converged'].strip().lower() not in ('yes', '1', 'true')]
print(f"graded rows with opt_converged != yes: {len(unconv_graded)}")

evals = [int(r['opt_evals_used']) for r in graded]
print(f"evals used: min={min(evals)} max={max(evals)} median={statistics.median(evals)} mean={statistics.mean(evals):.2f}")

# hand version columns
hv = set(r['hand_target_version'] for r in per)
ov_graded = set(r['opt_target_version'] for r in graded)
ov_rest = set(r['opt_target_version'] for r in per if r not in graded)
print(f"hand_target_version values: {hv}")
print(f"opt_target_version graded: {ov_graded}; rest: {ov_rest}")

# ---------- 2. regraded survey CSV ----------
with open(REG, newline='', encoding='utf-8') as f:
    reg = list(csv.DictReader(f))
print(f"\nregraded rows: {len(reg)}")
reg_built = [r for r in reg if r['built'].strip().lower() == 'true']
reg_cannot = [r for r in reg if r['built'].strip().lower() == 'false']
reg_graded = [r for r in reg if r['meetPct'].strip() not in ('', '0') or r['note'].strip() == '' or True]
# graded = numeric meetPct (>0 chars and parseable) AND note not GRADE-UNABLE/CANNOT BUILD
reg_graded = [r for r in reg if r['built'].strip().lower() == 'true' and 'GRADE-UNABLE' not in r['note']]
reg_meet_avg = sum(float(r['meetPct']) for r in reg_graded) / len(reg_graded)
reg_act_avg = sum(float(r['activePct']) for r in reg_graded) / len(reg_graded)
print(f"regraded: built={len(reg_built)} cannot={len(reg_cannot)} graded={len(reg_graded)}")
print(f"regraded avgs: meet={reg_meet_avg:.6f} active={reg_act_avg:.6f}")
check(len(reg) == 27, f"regraded CSV has 27 rows (got {len(reg)})")
check(len(reg_built) == 23, f"regraded built==23 (got {len(reg_built)})")
check(len(reg_graded) == 16, f"regraded graded==16 (got {len(reg_graded)})")
check(abs(reg_meet_avg - opt_meet_avg) < 1e-9 and abs(reg_act_avg - opt_act_avg) < 1e-9,
      f"opt averages from per-muscle == from regraded CSV ({opt_meet_avg:.4f}/{opt_act_avg:.4f} vs {reg_meet_avg:.4f}/{reg_act_avg:.4f})")

unconv = [r for r in reg_graded if str(r['converged']).strip() not in ('1', '1.0', 'true')]
print(f"regraded graded rows converged!=1: {len(unconv)}")

# per-muscle opt cells vs regraded, matched by muscle name
reg_by_name = {r['muscle']: r for r in reg}
mism = []
for r in graded:
    g = reg_by_name.get(r['muscle'])
    if g is None:
        mism.append(f"{r['muscle']}: no regraded row")
        continue
    if abs(num(r['opt_meet_pct']) - float(g['meetPct'])) > 1e-9: mism.append(f"{r['muscle']} meet {r['opt_meet_pct']} vs {g['meetPct']}")
    if abs(num(r['opt_active_pct']) - float(g['activePct'])) > 1e-9: mism.append(f"{r['muscle']} active {r['opt_active_pct']} vs {g['activePct']}")
    if r['opt_nBPA'] != g['nBPA']: mism.append(f"{r['muscle']} nBPA {r['opt_nBPA']} vs {g['nBPA']}")
    if r['opt_G'] != g['tackleG']: mism.append(f"{r['muscle']} G {r['opt_G']} vs {g['tackleG']}")
    if r['opt_evals_used'] != g['evalsUsed']: mism.append(f"{r['muscle']} evals {r['opt_evals_used']} vs {g['evalsUsed']}")
check(not mism, f"all 16 graded opt cells (meet/active/nBPA/G/evals) match regraded CSV{'' if not mism else '; ' + '; '.join(mism)}")

# ---------- 3. hand-tuned metrics CSV ----------
with open(HTM, newline='', encoding='utf-8') as f:
    htm = list(csv.DictReader(f))
print(f"\nhand-tuned metrics rows: {len(htm)}")
htm_by_name = {r['muscle']: r for r in htm}
mism = []
for r in per:
    h = htm_by_name.get(r['muscle'])
    if h is None: mism.append(f"{r['muscle']}: no hand-metrics row"); continue
    if abs(num(r['hand_meet_pct']) - float(h['meetPct'])) > 1e-9: mism.append(f"{r['muscle']} meet {r['hand_meet_pct']} vs {h['meetPct']}")
    if abs(num(r['hand_active_pct']) - float(h['activePct'])) > 1e-9: mism.append(f"{r['muscle']} active {r['hand_active_pct']} vs {h['activePct']}")
check(len(htm) == 27 and not mism, f"all 27 hand cells match hand_tuned_metrics CSV{'' if not mism else '; ' + '; '.join(mism)}")
htm_sus = [r['muscle'] for r in htm if r['suspect'].strip().lower() == 'yes']
check(sorted(htm_sus) == sorted(r['muscle'] for r in suspects),
      f"suspect sets identical between per-muscle and hand metrics ({len(suspects)} vs {len(htm_sus)})")

# ---------- 4. human torque curves CSV artifact flags ----------
with open(HTC, newline='', encoding='utf-8') as f:
    lines = [l for l in f if not l.startswith('#')]
htc = list(csv.DictReader(lines))
flagged = [(r['actuator_id'], r['master_name']) for r in htc if r['artifact_flag_caveat'].strip().lower() != 'none']
print(f"\nhuman curves rows: {len(htc)}; artifact-flagged: {flagged}")
check(len(htc) == 27, f"human curves CSV has 27 actuator rows (got {len(htc)})")
check(len(flagged) == 6, f"6 artifact-flagged curves (got {len(flagged)})")
art_names = set(r['muscle'] for r in artifacts)
flag_names = set(n for _, n in flagged)
check(art_names == flag_names, f"per-muscle path_artifact set == curves flagged set ({sorted(art_names)} vs {sorted(flag_names)})")

# ---------- 5. deliverable cells: md, tex, xlsx ----------
md = MD.read_text(encoding='utf-8')
m = re.search(r'\|\s*Hand-tuned ([^|]*)\|\s*([\d.]+)%\s*\|\s*([\d.]+)%\s*\|', md)
mo = re.search(r'\|\s*Optimized ([^|]*)\|\s*([\d.]+)%\s*\|\s*([\d.]+)%\s*\|', md)
md_cells = dict(hand_label=m.group(1).strip(), hand_meet=m.group(2), hand_act=m.group(3),
                opt_label=mo.group(1).strip(), opt_meet=mo.group(2), opt_act=mo.group(3))
print(f"\nmd cells: {md_cells}")

tex = TEX.read_text(encoding='utf-8')
t = re.search(r'Hand-tuned ([^&]*)&\s*([\d.]+)\s*\\%\s*&\s*([\d.]+)\s*\\%', tex)
to = re.search(r'Optimized ([^&]*)&\s*([\d.]+)\s*\\%\s*&\s*([\d.]+)\s*\\%', tex)
tex_cells = dict(hand_label=t.group(1).strip(), hand_meet=t.group(2), hand_act=t.group(3),
                 opt_label=to.group(1).strip(), opt_meet=to.group(2), opt_act=to.group(3))
print(f"tex cells: {tex_cells}")

wb = openpyxl.load_workbook(XLSX, data_only=True)
print(f"xlsx sheets: {wb.sheetnames}")
ws = wb[wb.sheetnames[0]]
xlsx_rows = []
for row in ws.iter_rows(values_only=True):
    xlsx_rows.append(row)
    print("  xlsx row:", row)
# find the two TABLE data rows: label contains the parenthesized coverage label
xl = {}
for row in xlsx_rows:
    cells = [str(c) if c is not None else '' for c in row]
    if cells and 'Hand-tuned (' in cells[0]: xl['hand'] = row
    if cells and 'Optimized (' in cells[0]: xl['opt'] = row
print(f"xlsx hand table row: {xl.get('hand')}")
print(f"xlsx opt table row:  {xl.get('opt')}")
# number formats of the pct cells
for key in ('hand', 'opt'):
    row = xl[key]
    fmts = [ws.cell(row=xlsx_rows.index(row) + 1, column=i + 1).number_format for i in range(1, 4)]
    print(f"xlsx {key} number formats (cols 2-4): {fmts}")

xl_hand = [c for c in xl['hand'][1:3]]
xl_opt = [c for c in xl['opt'][1:3]]
xl_hand_avgd = xl['hand'][3]
xl_opt_avgd = xl['opt'][3]
print(f"xlsx hand pct values: {xl_hand} (muscles averaged col: {xl_hand_avgd})")
print(f"xlsx opt pct values:  {xl_opt} (muscles averaged col: {xl_opt_avgd})")

# traceability sums in xlsx
tr = {}
for row in xlsx_rows:
    c0 = str(row[0]) if row[0] is not None else ''
    if 'sum / count' in c0: tr[c0.strip()] = (row[1], row[2])
print(f"xlsx traceability sums: {tr}")
hand_meet_sum = sum(m for m, _ in hand); hand_act_sum = sum(a for _, a in hand)
opt_meet_sum = sum(num(r['opt_meet_pct']) for r in graded); opt_act_sum = sum(num(r['opt_active_pct']) for r in graded)
check(tr.get('Hand-tuned meet sum / count') == (hand_meet_sum, 27), f"xlsx hand meet sum/count {tr.get('Hand-tuned meet sum / count')} == {hand_meet_sum}/27")
check(tr.get('Hand-tuned active sum / count') == (hand_act_sum, 27), f"xlsx hand active sum/count {tr.get('Hand-tuned active sum / count')} == {hand_act_sum}/27")
check(tr.get('Optimized meet sum / count') == (opt_meet_sum, 16), f"xlsx opt meet sum/count {tr.get('Optimized meet sum / count')} == {opt_meet_sum}/16")
check(tr.get('Optimized active sum / count') == (opt_act_sum, 16), f"xlsx opt active sum/count {tr.get('Optimized active sum / count')} == {opt_act_sum}/16")

R = {'hand_meet': f"{hand_meet_avg:.2f}", 'hand_act': f"{hand_act_avg:.2f}",
     'opt_meet': f"{opt_meet_avg:.2f}", 'opt_act': f"{opt_act_avg:.2f}"}
print(f"recomputed rounded: {R}")
check(md_cells['hand_meet'] == R['hand_meet'] and md_cells['hand_act'] == R['hand_act'], f"md hand cells == recomputed ({md_cells['hand_meet']}/{md_cells['hand_act']} vs {R['hand_meet']}/{R['hand_act']})")
check(md_cells['opt_meet'] == R['opt_meet'] and md_cells['opt_act'] == R['opt_act'], f"md opt cells == recomputed ({md_cells['opt_meet']}/{md_cells['opt_act']} vs {R['opt_meet']}/{R['opt_act']})")
check(tex_cells['hand_meet'] == R['hand_meet'] and tex_cells['hand_act'] == R['hand_act'], f"tex hand cells == recomputed ({tex_cells['hand_meet']}/{tex_cells['hand_act']})")
check(tex_cells['opt_meet'] == R['opt_meet'] and tex_cells['opt_act'] == R['opt_act'], f"tex opt cells == recomputed ({tex_cells['opt_meet']}/{tex_cells['opt_act']})")
check(len(xl_hand) == 2 and abs(float(xl_hand[0]) - hand_meet_avg) < 0.005 and abs(float(xl_hand[1]) - hand_act_avg) < 0.005,
      f"xlsx hand cells == recomputed ({xl_hand} vs {hand_meet_avg:.4f}/{hand_act_avg:.4f})")
check(len(xl_opt) == 2 and abs(float(xl_opt[0]) - opt_meet_avg) < 0.005 and abs(float(xl_opt[1]) - opt_act_avg) < 0.005,
      f"xlsx opt cells == recomputed ({xl_opt} vs {opt_meet_avg:.4f}/{opt_act_avg:.4f})")

# ---------- 6. row labels ----------
print(f"\nmd hand label: '{md_cells['hand_label']}'  tex: '{tex_cells['hand_label']}'")
print(f"md opt label:  '{md_cells['opt_label']}'  tex: '{tex_cells['opt_label']}'")
check('27 of 27' in md_cells['hand_label'], "md hand label says 27 of 27")
check('27 of 27' in tex_cells['hand_label'], "tex hand label says 27 of 27")
check('16 of 27 graded' in md_cells['opt_label'] and '23 built' in md_cells['opt_label'], "md opt label says 16 of 27 graded; 23 built")
check('16 of 27 graded' in tex_cells['opt_label'] and '23 built' in tex_cells['opt_label'], "tex opt label says 16 of 27 graded; 23 built")
check(len(graded) == 16 and built == 23, f"label counts match recomputation (graded={len(graded)}, built={built})")
check('16 of 27 graded' in str(xl['opt'][0]) and '23 built' in str(xl['opt'][0]), f"xlsx opt row label carries counts: '{xl['opt'][0]}'")
check('27 of 27' in str(xl['hand'][0]), f"xlsx hand row label carries counts: '{xl['hand'][0]}'")
check(xl_hand_avgd == 27 and xl_opt_avgd == 16, f"xlsx 'Muscles averaged' column 27/16 (got {xl_hand_avgd}/{xl_opt_avgd})")

# xlsx per-muscle sheet vs per-muscle CSV (key numeric columns)
ws2 = wb['Per-muscle rows']
hdr = [str(c.value) if c.value is not None else '' for c in next(ws2.iter_rows(max_row=1))]
print(f"\nxlsx per-muscle sheet header: {hdr}")
col = {name: i for i, name in enumerate(hdr)}
sheet_rows = []
for row in ws2.iter_rows(min_row=2, values_only=True):
    if row[0] is None: continue
    sheet_rows.append(row)
print(f"xlsx per-muscle data rows: {len(sheet_rows)}")
def sval(v):
    if v is None: return ''
    if isinstance(v, float): return f"{v:.4f}"
    return str(v).strip()
mism = []
for r_csv, r_xl in zip(per, sheet_rows):
    if r_csv['muscle'].strip() != str(r_xl[col.get('muscle', 0)]).strip():
        mism.append(f"name {r_csv['muscle']} vs {r_xl[col.get('muscle', 0)]}")
    for csv_col, xl_col in [('hand_meet_pct', 'hand_meet_pct'), ('hand_active_pct', 'hand_active_pct'),
                            ('opt_meet_pct', 'opt_meet_pct'), ('opt_active_pct', 'opt_active_pct')]:
        a, b = r_csv[csv_col].strip(), sval(r_xl[col.get(xl_col, -1)] if xl_col in col else '')
        if a == '':
            if b != '': mism.append(f"{r_csv['muscle']} {csv_col}: csv blank vs xlsx {b}")
        elif b == '' or abs(float(a) - float(b)) > 1e-9:
            mism.append(f"{r_csv['muscle']} {csv_col}: {a} vs {b}")
check(len(sheet_rows) == 27 and not mism,
      f"xlsx per-muscle sheet (27 rows) matches CSV on meet/active columns{'' if not mism else '; ' + '; '.join(mism[:10])}")

# ---------- 7. assemble summary ----------
S = json.loads(SUMJ.read_text(encoding='utf-8'))
print(f"\nassemble_summary: {S}")
check(S['hand_meet_avg'] == float(R['hand_meet']) and S['hand_act_avg'] == float(R['hand_act']), "summary hand avgs match")
check(S['opt_meet_avg'] == float(R['opt_meet']) and S['opt_act_avg'] == float(R['opt_act']), "summary opt avgs match")
check(S['hand_covered'] == 27 and S['opt_covered_graded'] == 16 and S['opt_built'] == 23, "summary counts 27/16/23 match")
check(S['could_not_build'] == len(cannot_build) == 4, "summary could_not_build==4 matches")
check(S['built_grade_unable'] == len(grade_unable) == 7, "summary built_grade_unable==7 matches")
check(S['transcription_suspects'] == len(suspects) == len(htm_sus), f"summary suspects=={len(suspects)} matches CSVs")
check(S['artifact_flagged_curves'] == len(flagged) == len(artifacts) == 6, "summary artifact_flagged==6 matches")
check(S['unconverged_surveyed_sites'] == 0 and len(unconv) == 0 and len(unconv_graded) == 0, "summary unconverged==0 matches (no converged!=1 rows)")

# ---------- 8. caveats-report claims vs data ----------
cav = (FOLDER / "CAVEATS_REPORT_20261022.md").read_text(encoding='utf-8')
for name, cnt in [("Could not build (4)", 4), ("Built but grade-unable (7)", 7)]:
    pass
check(f"Could not build (4)" in cav, "caveats states Could not build (4)")
check(f"Built but grade-unable (7)" in cav, "caveats states Built but grade-unable (7)")
check("23 of 27 built" in cav and "16 built muscles" in cav, "caveats states 23 of 27 built / 16 graded")
for nm in [r['muscle'] for r in cannot_build]:
    check(nm in cav, f"caveats lists cannot-build muscle {nm}")
for nm in [r['muscle'] for r in grade_unable]:
    check(nm in cav, f"caveats lists grade-unable muscle {nm}")
cav_suspect_names = re.findall(r'\*\*(.{2,40}?)\*\* \(act (\d+), hand meet ([\d.]+)%\)', cav)
print(f"caveats suspect bullets: {cav_suspect_names}")
check(len(cav_suspect_names) == 9, f"caveats lists 9 suspect bullets (got {len(cav_suspect_names)})")
per_by_id = {int(r['actuator_id']): r for r in per}
for nm, aid, hm in cav_suspect_names:
    r = per_by_id.get(int(aid))
    okname = r is not None and r['muscle'].strip() == nm.strip()
    okval = r is not None and abs(num(r['hand_meet_pct']) - float(hm)) < 1e-9
    check(okname and okval, f"caveats suspect '{nm}' act {aid} hand meet {hm} matches per-muscle row")
check(len([s for s in cav_suspect_names if int(s[1]) in [int(r['actuator_id']) for r in suspects]]) == 9,
      "caveats suspect ids == per-muscle suspect flag set")
for site in ['semimem_r', 'rect_fem_r', 'grac_r', 'vas_int_r', 'iliacus_r', 'psoas_r']:
    check(site in cav, f"caveats artifact section names {site}")
check("11 of 16" in cav and len(capped) == 11, f"caveats '11 of 16 capped' matches recomputed {len(capped)}")
check(f"34-195" in cav and min(evals) == 34 and max(evals) == 195, "caveats evals range 34-195 matches")
med_claim = re.search(r'median (\d+)', cav)
print(f"caveats median claim: {med_claim.group(1) if med_claim else None}; recomputed median of 16 evals = {statistics.median(evals)}")
check(med_claim is not None and abs(float(med_claim.group(1)) - statistics.median(evals)) < 1e-9,
      f"caveats median claim ({med_claim.group(1) if med_claim else '?'}) == recomputed ({statistics.median(evals)})")
check("meet 100.00% with active 0.00%" in cav.replace('\u2019', "'") or "100.00% with active 0.00%" in cav,
      "caveats flags the Bicep Femoris LH meet100/active0 suspect row")
bf = [r for r in per if 'Long Head' in r['muscle']][0]
check(num(bf['opt_meet_pct']) == 100.0 and num(bf['opt_active_pct']) == 0.0, "Bicep Femoris LH opt row is indeed 100/0")

print("\n" + "=" * 60)
print(f"TOTAL PROBLEMS: {len(problems)}")
for p in problems:
    print("  - " + p)
