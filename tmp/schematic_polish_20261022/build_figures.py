"""Editable dissertation schematics, sized in points at their printed width.

PowerPoint is the editable source; PDF and EPS exports are derived artifacts.
No experimental data are generated or changed by this module.
"""
from pathlib import Path
import math, json
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from PIL import ImageFont

ROOT=Path(__file__).resolve()
while not (ROOT/'Documentation').is_dir(): ROOT=ROOT.parent
FIG=ROOT/'Documentation/Reports and Papers/Dissertation/Figures'
OUT=ROOT/'tmp/schematic_polish_20261022/staged'
INK='243343'; BLUE='0072B2'; TEAL='009E73'; ORANGE='D55E00'; GOLD='E69F00'
GRAY='6C7680'; LIGHT='F1F5F8'; WHITE='FFFFFF'; PALE='E5F1F8'
JOBS=[]

def wrap(text,width,size=11,bold=False):
    font=ImageFont.truetype(r'C:\Windows\Fonts\arialbd.ttf' if bold else r'C:\Windows\Fonts\arial.ttf',round(size*20))
    lines=[]
    for paragraph in text.split('\n'):
        words=paragraph.split();current=''
        for word in words:
            proposed=(current+' '+word).strip()
            if current and font.getlength(proposed)/20>width-2:
                lines.append(current);current=word
            else:current=proposed
        lines.append(current)
    return '\n'.join(lines)

def rgb(v): return RGBColor.from_string(v)
def new(h):
    p=Presentation();p.slide_width=Pt(468);p.slide_height=Pt(h)
    return p,p.slides.add_slide(p.slide_layouts[6])
def line(s,x1,y1,x2,y2,color=INK,dash=False,width=1,arrow=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2))
    sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    ln=sh.line._get_or_add_ln()
    if dash:
        el=ln.makeelement(qn('a:prstDash'),{'val':'dash'});ln.append(el)
    if arrow:
        el=ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'sm','len':'sm'});ln.append(el)
    return sh
def path(s,pts,**kw):
    for i,(a,b) in enumerate(zip(pts,pts[1:])):
        line(s,*a,*b,arrow=(i==len(pts)-2),**kw)
def shape(s,x,y,w,h,kind=MSO_SHAPE.RECTANGLE,fill=WHITE,dash=False,color=INK,width=1):
    sh=s.shapes.add_shape(kind,Pt(x),Pt(y),Pt(w),Pt(h))
    if fill: sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    else: sh.fill.background()
    sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    if dash:sh.line._get_or_add_ln().append(sh.line._get_or_add_ln().makeelement(qn('a:prstDash'),{'val':'dash'}))
    return sh
def txt(s,x,y,w,h,text,size=11,bold=False,align='left',color=INK,margin=0):
    sh=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h))
    tf=sh.text_frame;tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=Pt(margin);tf.margin_top=tf.margin_bottom=Pt(margin)
    text=wrap(text,w-margin*2,size,bold)
    for i,t in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(size*1.10)
        r=p.add_run();r.text=t;r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.italic=False;r.font.color.rgb=rgb(color)
    return sh
def box(s,x,y,w,h,title,sub=None,dash=False,fill=LIGHT,size=11):
    sh=shape(s,x,y,w,h,fill=fill,dash=dash)
    title=wrap(title,w-14,size,sub is not None)
    if sub is not None:sub=wrap(sub,w-14,size)
    if sub is None:
        txt(s,x+7,y+(h-len(title.split('\n'))*size*1.10)/2,w-14,h-8,title,size=size,align='center')
    else:
        n=len(title.split('\n'));m=len(sub.split('\n'));height=(n+m)*size*1.10+5
        top=y+(h-height)/2+.5
        txt(s,x+7,top,w-14,n*size*1.15,title,size=size,bold=True,align='center')
        txt(s,x+7,top+n*size*1.10+5,w-14,m*size*1.16,sub,size=size,align='center')
    return sh
def tag(s,x,y,text,w=100):txt(s,x,y,w,28,text,size=10.8,align='center',color=GRAY)
def save(p,chapter,name,alt):
    d=OUT/chapter/'components';d.mkdir(parents=True,exist_ok=True)
    p.save(d/(name+'.pptx'));(d/(name+'.alt.txt')).write_text(alt,encoding='utf-8')
    JOBS.append((chapter,name))

def roadmap():
    p,s=new(242)
    line(s,15,12,36,12);txt(s,43,5,214,16,'Developed or used in this dissertation',size=10.8)
    line(s,295,12,316,12,dash=True);txt(s,323,5,130,16,'Future integration',size=10.8)
    xs=[28,178,328];w=112;h=51;yt=53;yb=153
    data=[('Biomechanics','Human benchmark\n(Ch. 2)'),('BPA force models','Characterization\n(Ch. 3)'),('Torque model','Corrections, routes\n(Ch. 4)'),('Sensory Afferent\nDatabase','Pathways (Ch. 5)'),('SNS simulation','Software tools\n(Ch. 6)'),('BPA robot','SNS control\n(future work)')]
    for i,(a,b) in enumerate(data):box(s,xs[i%3],yt if i<3 else yb,w,h,a,b,dash=i==5,size=10.8)
    for x in xs[:2]:line(s,x+w,78.5,x+150,78.5,arrow=True)
    tag(s,126,30,'targets',66);tag(s,274,23,'actuator model',66)
    line(s,140,178.5,178,178.5,arrow=True);line(s,290,178.5,328,178.5,dash=True,arrow=True)
    path(s,[(384,104),(384,126),(234,126),(234,153)])
    txt(s,255,131,179,14,'Routes and torque model',size=10.8,align='right')
    path(s,[(384,204),(384,225),(12,225),(12,78.5),(28,78.5)],dash=True)
    txt(s,149,207,170,15,'Biological insight (future)',size=10.8,align='center')
    save(p,'10-introduction','roadmap','Chapter roadmap linking biomechanics, force characterization, torque corrections, sensory evidence, neuromechanical software, and future robot integration. Dashed lines distinguish future integration and its biological-insight return path.')

def pipeline():
    p,s=new(303)
    for x,w,t in [(10,128,'Model conversion'),(160,150,'Closed-loop simulation'),(332,126,'Future hardware')]:
        txt(s,x,7,w,18,t,size=12,bold=True,align='center')
        if x==332:shape(s,x,31,w,222,fill=None,dash=True,color=GRAY)
    box(s,16,40,116,48,'OpenSim model','Gait2392 routing')
    box(s,16,113,116,34,'MyoConverter')
    box(s,16,174,116,51,'MuJoCo model','MJCF; activation\nstate actuators')
    line(s,74,88,74,113,arrow=True);line(s,74,147,74,174,arrow=True)
    box(s,183,40,105,58,'MuJoCo physics','BPA force model\n(Ch. 3)')
    box(s,183,187,105,42,'SNS-Toolbox','Neural controller')
    path(s,[(132,199),(150,199),(150,68),(183,68)])
    path(s,[(183,208),(169,208),(169,113),(209,113),(209,98)],color=BLUE)
    txt(s,180,123,109,26,'Per-muscle\nactivation',size=10.8,align='left',color=BLUE)
    path(s,[(271,98),(271,112),(305,112),(305,180),(262,180),(262,187)],color=INK)
    txt(s,184,153,110,28,'Ia, Ib, II, and\ncontact feedback',size=10.8)
    box(s,343,40,103,45,'Jetson Orin','SNS evaluation',dash=True)
    box(s,343,122,103,50,'Joint controllers','Valves, pressure,\nand contact',dash=True)
    box(s,343,210,103,30,'Treadmill biped',dash=True)
    line(s,395,85,395,122,arrow=True);line(s,395,172,395,210,arrow=True)
    txt(s,343,90,42,26,'SNS\noutput',size=10.8,align='center')
    txt(s,340,178,46,26,'Valve\ndrive',size=10.8,align='center')
    path(s,[(446,225),(452,225),(452,146),(446,146)],color=GRAY)
    path(s,[(343,136),(335,136),(335,63),(343,63)],dash=True,color=GRAY)
    path(s,[(288,208),(322,208),(322,62),(343,62)],dash=True)
    txt(s,13,268,442,29,'Computer-in-the-loop: test embedded control against the simulated body\nbefore transfer to the physical robot.',size=10.8,align='center')
    save(p,'60-sns','pipeline','Conversion from OpenSim through MyoConverter to MuJoCo, a closed loop between the BPA-actuated physics and SNS controller, and a dashed future embedded-hardware path. Per-muscle activation and proprioceptive/contact feedback run in opposite directions.')

def curriculum():
    p,s=new(287)
    cols=[(14,27,'No.'),(48,110,'Preparation'),(176,128,'Capability'),(324,130,'Score')]
    for x,w,t in cols:txt(s,x,8,w,18,t,bold=True,size=11)
    line(s,12,31,456,31,color=GRAY)
    rows=[('1','Deafferented\nair stepping','Intrinsic rhythm;\nfeedback off','Rhythm and swing\n(rhythm-gated)'),('2','Afferented\nair stepping','Heel, toe, and load\nfeedback active','Rhythm and swing\n(rhythm-gated)'),('3','Standing\nbalance','Posture under\npush perturbations','Posture maintenance\nunder push'),('4','Ground\nwalking','Ground contact\nand forward gait','Kinematics vs. human\nreference cycle'),('5','Contact-driven\nground walking','Contact-driven\npathway variant','Kinematics vs. human\nreference cycle')]
    for i,(n,a,b,c) in enumerate(rows):
        y=40+i*41
        txt(s,16,y,24,28,n,size=12,bold=True,color=BLUE)
        txt(s,48,y,118,30,a,bold=True);txt(s,176,y,135,30,b);txt(s,324,y,132,30,c)
        line(s,12,y+34,456,y+34,color='D9E1E7',width=.6)
    txt(s,14,252,440,29,'Each winner seeds the next stage. Finalists undergo perturbation and basin checks.\nAir stages are suspended; ground evaluations retain the support harness.',size=10.8)
    save(p,'60-sns','sns_curriculum','Five curriculum stages with the capability and score of each. Intrinsic air-stepping rhythm is tested without feedback, followed by sensory air stepping, standing, ground walking, and the contact-driven variant. Ground evaluations retain a support harness; stage number does not imply its withdrawal.')

def library():
    p,s=new(211)
    xs=[12,170,328];ws=128
    entries=[('Graded neuron','Membrane voltage\nis the output'),('Graded synapse','Conductance coupling;\nexcitation or inhibition'),('Spiking neuron','Leaky integrate-\nand-fire model'),('Ia encoder','Muscle length\nand velocity'),('Ib encoder','Muscle force'),('Activation dynamics','Continuous muscle\nactivation filter')]
    for i,(a,b) in enumerate(entries):box(s,xs[i%3],12+(i//3)*85,ws,65,a,b)
    line(s,12,178,456,178,color=GRAY)
    txt(s,14,187,440,18,'Also: synapse sum, diameter-specific BPA force blocks, and biological muscle.',size=10.8,align='center')
    save(p,'60-sns','sns_library_overview','Six implemented SNS library component types, arranged by neural components and sensory/activation components, with synapse-sum and force-model additions listed below. This inventory is distinct from the provisional generic BPA block.')

def workflow():
    p,s=new(328)
    box(s,12,13,130,55,'Human benchmark','Torque and range\nof motion (Ch. 2)')
    box(s,12,117,130,50,'BPA force model','Actuator forces\n(Ch. 3)')
    box(s,196,13,145,55,'Select actuators','Diameter, resting length,\nand actuator count')
    box(s,196,117,145,55,'Design routes','Corrected torque model\n(Ch. 4)')
    line(s,142,40,196,40,arrow=True);line(s,268.5,68,268.5,117,arrow=True)
    line(s,142,139,196,139,arrow=True)
    tag(s,143,18,'targets',50)
    shape(s,212,207,113,63,kind=MSO_SHAPE.DIAMOND,fill=WHITE)
    txt(s,231,226,75,28,'Meets\ntargets?',size=11,align='center')
    line(s,268.5,172,268.5,207,arrow=True)
    txt(s,281,179,164,17,'Predicted torque and motion',size=10.8)
    box(s,17,215,145,50,'Refine design','Redesign routes or\nreselect actuators')
    line(s,212,238.5,162,238.5,arrow=True);tag(s,166,214,'no',40)
    line(s,89.5,215,89.5,185);line(s,89.5,185,178,185)
    line(s,178,185,178,147);line(s,178,131,178,40)
    line(s,178,40,196,40,arrow=True)
    box(s,355,214,101,55,'Closed-loop\nverification\n(Ch. 6)',dash=True,size=10.8)
    line(s,325,238.5,355,238.5,arrow=True);tag(s,322,214,'yes',34)
    box(s,342,291,114,28,'Commit to hardware',dash=True,size=10.8)
    line(s,405,269,405,291,arrow=True)
    save(p,'70-overall','design_workflow','Actuator selection and corrected route design feed a torque-and-motion target check. A failed design returns to route or actuator selection. Passing designs proceed to closed-loop simulation verification before hardware commitment.')

def sensorium():
    p,s=new(347)
    txt(s,12,6,444,18,'Proposed supraspinal additions',bold=True,size=12)
    for x,a,b in [(12,'Basal ganglia','Action and mode\nselection'),(170,'Cerebellum','Adaptive gait gating\nand phase tuning'),(328,'Vestibular / ocular','Orientation sensing\nand visual correction')]:box(s,x,34,128,58,a,b,dash=True)
    txt(s,12,127,410,18,'Spinal network (Ch. 6)',bold=True,size=12)
    box(s,12,153,174,50,'Two-level CPG','Rhythm and pattern formation')
    box(s,237,153,174,50,'Sensorimotor layer','Motor pools; Ia, Ib, II;\nheel and toe contact')
    line(s,186,178,237,178,arrow=True)
    path(s,[(237,192),(218,192),(218,235),(99,235),(99,203)])
    txt(s,107,214,105,17,'Sensory feedback',size=10.8)
    path(s,[(76,92),(76,106),(174,106),(174,153)],dash=True)
    path(s,[(234,92),(234,117),(220,117),(220,145),(174,145),(174,153)],dash=True)
    path(s,[(392,92),(392,116),(367,116),(367,153)],dash=True)
    txt(s,254,101,100,29,'Descending\ncommands',size=10.8)
    txt(s,12,249,190,17,'Simulated mechanical model',bold=True,size=12)
    box(s,12,274,399,32,'BPA actuators and robot model')
    line(s,286,203,286,274,arrow=True)
    line(s,351,274,351,203,arrow=True)
    txt(s,211,216,68,16,'activation',size=10.8,align='right')
    txt(s,358,214,99,29,'Proprioceptive\nand contact input',size=10.8)
    path(s,[(411,290),(448,290),(448,63),(456,63)],dash=True,color=GRAY)
    txt(s,13,322,430,16,'Balance-platform hardware is a precursor to orientation feedback.',size=10.8,color=GRAY)
    save(p,'70-overall','sensorium_hierarchy','Proposed basal-ganglia, cerebellar, and vestibular/ocular modules above the implemented spinal and simulated BPA-body tiers. Descending, activation, proprioceptive/contact, and proposed orientation feedback paths are distinguished. A note identifies balance-platform hardware only as a precursor.')

def bpa():
    p,s=new(235)
    txt(s,12,5,444,18,'(A) Layered FEA model: loads and constraints',size=12,bold=True)
    x0,x1=72,313;yt=53;bh=15
    shape(s,53,45,19,116,fill='D9E1E7')
    for y in range(46,158,10):line(s,44,y+7,53,y,width=.7)
    txt(s,6,70,40,30,'Fixed\nend',size=11)
    names=['Fabric outer','Outer rubber','Aramid helix (+θ)','Middle rubber','Inner helix (−θ)','Inner rubber']
    colors=['D7DEE4','CCE5F3','F6D7C8','CCE5F3','FAEABB','CCE5F3']
    for i,(n,c) in enumerate(zip(names,colors)):
        y=yt+i*bh;shape(s,x0,y,x1-x0,bh,fill=c,color=GRAY,width=.55)
        txt(s,x1+15,y+1,130,14,n,size=10.8)
        line(s,x1,y+bh/2,x1+10,y+bh/2,width=.7)
    for i in range(5):
        y=yt+(i+1)*bh
        line(s,91,y-3,91,y+3,color=ORANGE,width=2)
    txt(s,82,27,220,17,'Contact elements at layer interfaces',size=10.8)
    line(s,91,45,91,51,color=ORANGE)
    for x in [120,165,210,255]:line(s,x,170,x,143,color=BLUE,arrow=True)
    txt(s,152,174,143,16,'Internal air pressure p',size=11,color=BLUE)
    line(s,62,200,313,200,dash=True,color=GRAY)
    txt(s,62,210,250,17,'Longitudinal symmetry plane',size=10.8)
    line(s,313,44,313,201,dash=True)
    txt(s,315,28,142,16,'Half-length boundary',size=10.8,bold=True)
    txt(s,335,161,119,14,'Radial motion allowed',size=10.8)
    line(s,327,171,327,190,arrow=True);line(s,327,171,327,151,arrow=True)
    txt(s,335,198,119,31,'Axial: fixed, or\nmovable / loaded',size=10.8)
    save(p,'70-overall','bpa_construction','Editable FEA wall-layer schematic with six layers, interlayer contact, outward internal pressure, fixed cap, longitudinal symmetry, and transverse half-length boundary. Radial motion remains free away from the cap. Axial displacement can be fixed, or a movable/load-controlled boundary can replace transverse symmetry.')

def actuator():
    p,s=new(321)
    panels=[(12,8,'(A) Electric torque motor'),(250,8,'(B) Braided pneumatic actuator'),(12,169,'(C) Hydraulic cylinder'),(250,169,'(D) Dielectric elastomer')]
    for x,y,t in panels:txt(s,x,y,206,18,t,size=11.5,bold=True)
    # Motor: stator, rotor, output shaft, rotational arrows.
    shape(s,55,42,83,83,kind=MSO_SHAPE.OVAL,fill=LIGHT)
    shape(s,78,65,37,37,kind=MSO_SHAPE.OVAL,fill='E8D7A2')
    for a in range(0,360,60):
        r=math.radians(a+30);shape(s,96.5+33*math.cos(r)-4,83.5+33*math.sin(r)-4,8,8,fill=BLUE,color=BLUE)
    line(s,96,83,195,83,width=3);path(s,[(60,37),(87,30),(117,34),(137,49)],color=BLUE)
    txt(s,17,137,203,17,'Electromagnetic shaft rotation',size=10.8,align='center')
    # BPA: explicitly outward wall pressure, distinct inward axial contraction.
    shape(s,284,62,136,43,fill=PALE)
    for x in range(287,419,13):
        if x+40<=418:
            line(s,x,64,x+40,103,color=ORANGE,width=.65)
            line(s,x,103,x+40,64,color=BLUE,width=.65)
    shape(s,277,59,7,49,fill=GRAY);shape(s,420,59,7,49,fill=GRAY)
    for x in [307,350,393]:line(s,x,77,x,49,color=TEAL,arrow=True);line(s,x,91,x,119,color=TEAL,arrow=True)
    line(s,259,83,277,83,arrow=True);line(s,446,83,427,83,arrow=True)
    txt(s,273,33,161,17,'Internal pressure',size=10.8,align='center',color=TEAL)
    txt(s,256,137,198,17,'Radial expansion; axial contraction',size=10.8,align='center')
    # Hydraulic piston.
    shape(s,30,222,122,51,fill='DDF0E9');shape(s,95,225,7,45,fill=GRAY)
    line(s,102,247,198,247,width=4);line(s,53,234,87,234,color=TEAL,arrow=True)
    line(s,157,228,199,228,color=BLUE,arrow=True)
    txt(s,17,290,202,16,'Fluid pressure produces linear force',size=10.8,align='center')
    # DEA: membrane and electrodes; compression inward, spreading outward.
    shape(s,293,231,126,38,fill='F1DFED',color=GRAY)
    line(s,293,230,419,230,color='AA4499',width=2);line(s,293,270,419,270,color='AA4499',width=2)
    line(s,356,204,356,225,color='AA4499',arrow=True);line(s,356,294,356,275,color='AA4499',arrow=True)
    line(s,294,250,271,250,color=BLUE,arrow=True);line(s,419,250,442,250,color=BLUE,arrow=True)
    txt(s,257,306,198,13,'Electric field: thinning and expansion',size=10.8,align='center')
    save(p,'15-background','actuator_technology_comparison','Four actuator mechanisms: electric motor rotation, BPA internal pressure acting outward with axial contraction, a pressure-driven hydraulic piston, and dielectric-elastomer thinning with lateral expansion. Direction and labels supplement the color coding.')

def reflex():
    p,s=new(344)
    # Explicit hierarchy; triangular excitatory and round inhibitory terminations.
    def node(x,y,t,kind=MSO_SHAPE.OVAL,w=29,h=29):
        shape(s,x,y,w,h,kind=kind,fill=PALE);txt(s,x,y+7,w,15,t,size=11,bold=True,align='center')
    def syn(pts,exc=True):
        for a,b in zip(pts,pts[1:]):line(s,*a,*b,width=1)
        x,y=pts[-1]
        if exc:
            sh=shape(s,x-5,y-4,7,8,kind=MSO_SHAPE.ISOSCELES_TRIANGLE,fill=WHITE,width=1)
            sh.rotation=90
        else:shape(s,x-3,y-3,6,6,kind=MSO_SHAPE.OVAL,fill=INK)
    txt(s,12,7,444,18,'Antagonist knee-reflex circuit',size=12,bold=True)
    txt(s,20,43,105,15,'Extensor sensors',size=10.8,bold=True)
    txt(s,26,145,105,15,'Flexor sensors',size=10.8,bold=True)
    txt(s,158,31,122,18,'Motor pools',bold=True);txt(s,303,31,145,18,'Actuation and joint',bold=True)
    for x,y,t in [(33,64,'Ia'),(33,107,'Ib'),(33,163,'Ia'),(33,206,'Ib')]:node(x,y,t)
    node(183,78,'MN');node(183,190,'MN')
    txt(s,164,112,68,15,'extensor',size=10.8,align='center');txt(s,164,225,68,15,'flexor',size=10.8,align='center')
    syn([(62,78),(173,86)],True);syn([(62,121),(173,100)],False)
    syn([(62,177),(173,198)],True);syn([(62,220),(173,212)],False)
    syn([(62,78),(96,78),(142,184),(173,205)],False)
    syn([(62,177),(110,177),(142,63),(173,93)],False)
    # Descending input uses separate branches above the motor column.
    txt(s,260,57,115,16,'Descending drive',size=10.8)
    path(s,[(251,65),(242,65),(242,77),(216,77)])
    line(s,242,77,242,94)
    line(s,242,107,242,186)
    line(s,242,186,216,186,arrow=True)
    box(s,280,91,82,40,'Activation\nand BPA',size=10.8)
    box(s,280,201,82,40,'Activation\nand BPA',size=10.8)
    line(s,212,93,280,111,arrow=True);line(s,212,205,280,221,arrow=True)
    box(s,396,138,60,49,'Knee','1 DoF',size=10.8)
    path(s,[(362,111),(378,111),(378,150),(396,150)])
    path(s,[(362,221),(378,221),(378,175),(396,175)])
    txt(s,365,91,94,15,'extensor torque',size=10.8)
    txt(s,280,248,100,15,'flexor torque',size=10.8,align='center')
    path(s,[(426,187),(426,282),(18,282),(18,78),(33,78)],dash=True,color=GRAY)
    for y in [121,177,220]:line(s,18,y,33,y,dash=True,color=GRAY,arrow=True)
    txt(s,111,290,330,15,'Mechanical feedback: length, velocity, and force',size=10.8,align='center')
    shape(s,22,319,8,8,kind=MSO_SHAPE.ISOSCELES_TRIANGLE,fill=WHITE).rotation=90
    txt(s,36,317,100,17,'Excitation',size=10.8)
    shape(s,151,320,6,6,kind=MSO_SHAPE.OVAL,fill=INK)
    txt(s,165,317,100,17,'Inhibition',size=10.8)
    line(s,280,323,303,323,dash=True,color=GRAY)
    txt(s,311,317,140,17,'Sensory feedback',size=10.8)
    save(p,'60-sns','KneeReflex_circuit','Antagonist knee reflex with Ia homonymous excitation and reciprocal inhibition, Ib homonymous inhibition, descending drive to both motor pools, activation and BPA force blocks, antagonist torques on a one-degree-of-freedom knee, and dashed mechanical sensory feedback. No additional interneuron populations are implied.')

def neuron_detail():
    p,s=new(270)
    box(s,12,10,110,55,'Synaptic sums','sumG = Σ gi\nsumE = Σ gi Ei',size=10.8)
    box(s,170,10,143,55,'Synaptic current','Isyn = sumE\n− V · sumG',size=10.8)
    box(s,350,10,106,55,'Leak current','Ileak = Gm\n· (Vrest − V)',size=10.8)
    line(s,122,37,170,37,arrow=True)
    path(s,[(241,65),(241,81),(50,81),(50,108)])
    path(s,[(403,65),(403,94),(125,94),(125,108)])
    box(s,12,108,145,55,'Current sum','Σ Iapp + Isyn\n+ Ileak',size=10.8)
    box(s,190,108,105,55,'Membrane gain','1000 / Cm',size=10.8)
    box(s,338,108,118,55,'Integrator 1/s','V_mV\nV(0) = Vrest',size=10.8)
    line(s,157,135.5,190,135.5,arrow=True);line(s,295,135.5,338,135.5,arrow=True)
    path(s,[(397,163),(397,181),(104,181),(104,205)])
    box(s,12,205,185,43,'Threshold and gain','(V − Thr) / Slope',size=10.8)
    box(s,250,205,112,43,'Saturation','Range: 0 to 1',size=10.8)
    line(s,197,226.5,250,226.5,arrow=True);line(s,362,226.5,388,226.5,arrow=True)
    txt(s,396,220,68,17,'S_drive',size=10.8)
    txt(s,12,254,444,15,'V is the membrane voltage fed back to both current calculations.',size=10.8,align='center')
    save(p,'60-sns','simulink_neuron_detail','Functional schematic transcribed from SNS_Library.slx graded-neuron implementation. Synaptic conductance and reversal-weighted sums determine synaptic current; leak and applied currents join it before membrane gain and integration. Voltage is fed back to both current calculations and transformed by threshold, inverse slope, and saturation into a bounded drive output.')

def bpa_cutaway():
    p,s=new(246)
    s.shapes.add_picture(str(FIG/'70-overall/components/bpa_exploded_cutaway_unlabeled.png'),Pt(12),Pt(50),width=Pt(330))
    txt(s,12,5,300,18,'(B) Exploded half-cutaway',size=12,bold=True)
    txt(s,12,27,105,17,'Fixed end cap',size=11)
    line(s,45,48,45,111,width=.8)
    targets=[((282,71),(344,47),'Fabric outer\ncovering',36),((307,95),(344,82),'Outer rubber',75),((305,118),(344,113),'Aramid helix (+θ)',106),((310,145),(344,144),'Aramid helix (−θ)',137),((311,153),(344,179),'Inner rubber',172)]
    for (x,y),(a,b),label,ty in targets:
        line(s,x,y,a,b,width=.8)
        shape(s,x-1.7,y-1.7,3.4,3.4,kind=MSO_SHAPE.OVAL,fill=INK,width=.5)
        txt(s,352,ty,108,28,label,size=10.8)
    line(s,250,166,250,202,width=.8)
    shape(s,248.3,164.3,3.4,3.4,kind=MSO_SHAPE.OVAL,fill=INK,width=.5)
    txt(s,155,209,190,28,'Longitudinal symmetry plane',size=10.8,align='center')
    save(p,'70-overall','bpa_exploded_cutaway','Illustrative exploded upper-half cutaway, with editable labels for fixed end cap, fabric covering, outer rubber, opposing aramid helices, inner rubber, and longitudinal symmetry plane. The plane label points to the translucent plane; the inner-rubber label points to retained material above it. Raster rendering is illustrative and the labeling remains editable in PowerPoint.')

if __name__=='__main__':
    for f in [roadmap,pipeline,curriculum,library,workflow,sensorium,bpa,actuator,reflex,neuron_detail,bpa_cutaway]:f()
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/'jobs.json').write_text(json.dumps(JOBS))
    print('Prepared',len(JOBS),'editable figures in',OUT)
