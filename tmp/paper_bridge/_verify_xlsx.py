import openpyxl, glob, os
base = r"D:/GitHub/Bipedal_Robot/Documentation/Reports and Papers/Knee_Torque_Test/Tables"
for name in ["PinnedFlexor_Preopt_GoF.xlsx", "PinnedFlexor_GoF.xlsx", "PinnedExtensor_GoF.xlsx",
             "AdoptedCorrectionTerms.xlsx", "Route_PinnedBrackets.xlsx", "Route_Biomimetic10mm.xlsx",
             "Route_Redesign20mm.xlsx"]:
    p = os.path.join(base, name)
    print("=" * 20, name)
    try:
        wb = openpyxl.load_workbook(p)
        for ws in wb.worksheets:
            print(f"-- sheet {ws.title} ({ws.max_row}x{ws.max_column})")
            for row in ws.iter_rows(max_row=min(ws.max_row, 14), values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    print(" | ".join(cells)[:200])
    except Exception as e:
        print("ERROR:", e)
