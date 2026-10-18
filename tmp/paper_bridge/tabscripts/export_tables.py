# -*- coding: utf-8 -*-
"""Export the five paper tables (plan.md section 5) to .xlsx workbooks in
Documentation/Reports and Papers/Knee_Torque_Test/Tables/ (ONLY .xlsx files
are created there; Ben's CoefficientTable.xlsx is untouched).

Every number comes from mining/results_ledger.md or the dissertation chapter
of record (30-torque.tex, committed 0d23da03):
  T1 PinnedFlexor_GoF.xlsx        chapter tab:FlxPinGof, 30-torque.tex:364-383 (= ledger section 3)
  T2 PinnedExtensor_GoF.xlsx      chapter tab:ExtPinGof, 30-torque.tex:400-420 (= ledger section 3)
  T3 AdoptedCorrectionTerms.xlsx  ledger section 1 (adopted values + search bounds)
  T4 RouteRedesign_Summary.xlsx   ledger sections 6.1-6.4 + chapter :504,:510,:512,:514,:516,:542
  T5 BiPulley_Batch.xlsx          chapter tab:bipulley_batch :520-534 + ledger section 6.5 (clearances)

The script then verifies itself: recomputes the T1/T2 mean rows from the
per-test cells (tolerance = printed-precision half-ulp, since the chapter
means were computed from full-precision values), re-reads every workbook,
and diffs the target directory before/after.

Run: D:/Anaconda/python.exe tmp/paper_bridge/tabscripts/export_tables.py
"""
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

OUT_DIR = os.path.join(
    r"D:\GitHub\Bipedal_Robot",
    "Documentation", "Reports and Papers", "Knee_Torque_Test", "Tables",
)

NM = "N\u00b7m"          # N·m
CHI = "\u03c7"           # χ

CAP_FONT = Font(italic=True, size=10)
HDR_FONT = Font(bold=True)
MEAN_FONT = Font(bold=True)
THIN = Side(style="thin", color="999999")
HDR_BORDER = Border(bottom=THIN)


def build(path, caption, header, rows, widths, mean_flags=None, num_formats=None):
    """rows: list of row tuples; mean_flags: set of row indices to bold;
    num_formats: optional {column_index_1based: excel format string} applied to
    numeric cells so the of-record precision is displayed (21.650, not 21.65)."""
    mean_flags = mean_flags or set()
    num_formats = num_formats or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Table"
    ncols = len(header)
    # Caption row (merged, wrapped, italic)
    ws.cell(row=1, column=1, value=caption).font = CAP_FONT
    ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.row_dimensions[1].height = 58
    # Header row
    for j, h in enumerate(header, start=1):
        c = ws.cell(row=2, column=j, value=h)
        c.font = HDR_FONT
        c.border = HDR_BORDER
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    # Data rows
    for i, row in enumerate(rows):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=3 + i, column=j, value=v)
            if i in mean_flags:
                c.font = MEAN_FONT
            if isinstance(v, float) and j in num_formats:
                c.number_format = num_formats[j]
            if isinstance(v, str) and len(v) > 40:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = "A3"
    wb.save(path)
    return path


# ----------------------------------------------------------------------
# T1 - Pinned-knee flexor goodness of fit (chapter tab:FlxPinGof)
# rows: (test, rmse_imp, fvu_imp, max_imp, rmse_cor, fvu_cor, max_cor)
# ----------------------------------------------------------------------
T1_ROWS = [
    ("l0 = 48.5 cm (training)",          7.617, 2.1278, 25.531, 2.165, 0.1718, 5.467),
    ("l0 = 45.7 cm (training)",          8.803, 1.9380, 26.264, 1.368, 0.0468, 3.614),
    ("l0 = 47.9 cm (validation)",        7.670, 1.8463, 21.650, 2.287, 0.1641, 7.818),
    ("l0 = 40.6 cm, tendon (training)",  6.373, 0.5935, 16.859, 1.611, 0.0427, 5.891),
    ("l0 = 41.0 cm (training)",          9.694, 2.1149, 21.076, 1.084, 0.0277, 1.766),
    ("Mean (training, 4 tests)",         8.122, 1.6935, 22.433, 1.557, 0.0723, 4.184),
    ("Mean (validation)",                7.670, 1.8463, 21.650, 2.287, 0.1641, 7.818),
]

# ----------------------------------------------------------------------
# T2 - Pinned-knee extensor goodness of fit (chapter tab:ExtPinGof)
# ----------------------------------------------------------------------
T2_ROWS = [
    ("l0 = 40.0 cm (fitting)",           3.092, 0.5638, 4.842, 1.244, 0.1303, 2.196),
    ("l0 = 40.0 cm, tendon (fitting)",   3.622, 6.3919, 7.845, 1.350, 0.8878, 1.890),
    ("l0 = 43.6 cm (fitting)",           2.437, 0.9618, 4.284, 0.978, 0.1634, 1.818),
    ("l0 = 43.2 cm, tendon (fitting)",   2.878, 4.7304, 4.817, 1.525, 1.3270, 2.481),
    ("l0 = 45.7 cm (validation)",        2.178, 1.0586, 4.974, 0.490, 0.0535, 1.206),
    ("l0 = 46.5 cm (fitting)",           0.949, 0.3152, 2.272, 0.699, 0.1714, 1.559),
    ("Mean (fitting, 5 tests)",          2.596, 2.5926, 4.812, 1.159, 0.5360, 1.989),
    ("Mean (validation)",                2.178, 1.0586, 4.974, 0.490, 0.0535, 1.206),
]

# ----------------------------------------------------------------------
# T3 - Adopted correction terms (ledger section 1)
# rows: (parameter, configuration, role, value_full, unit, value_reported,
#        bound_lo, bound_hi, bound_unit)
# ----------------------------------------------------------------------
T3_ROWS = [
    (CHI + "0", "Flexor", "Constant effective-length offset",
     3.9374, "mm", "+3.94 mm", 0, 20, "mm"),
    (CHI + "1", "Flexor", "Effective axial stiffness (bracket-frame x and z)",
     39978.7, "N/m", "3.998e4 N/m", 30000, 1000000, "N/m"),
    (CHI + "2", "Flexor", "Effective bending stiffness (bracket-frame y)",
     14733.8, "N/m", "1.473e4 N/m", 5000, 20000, "N/m"),
    (CHI + "0", "Extensor", "Constant effective-length offset (opposite sign to the flexor value)",
     -6.3780, "mm", "-6.38 mm", -20, 5, "mm"),
    (CHI + "3", "Extensor", "Wrapping-loss (fictive-strain) scale factor",
     0.158289, "unitless", "0.158", 0, 1, "unitless"),
]

# ----------------------------------------------------------------------
# T4 - Route-redesign summary (ledger 6.1-6.4 + chapter)
# ----------------------------------------------------------------------
T4_ROWS = [
    ("Direct extensor",
     "35.45 to 109.4", "20.25 to 50.32", 5.02, -20.2,
     "Meets the required envelope at every angle with zero shortfall points; "
     "superseded by the cam redesign."),
    ("Cam extensor (design of record)",
     "not reported", "20.25 to 50.32", 30.83, "not reported",
     "Current extensor design; pressure to meet the envelope spans 269 to 507 kPa; "
     "tibial wrap bend radius 96.9 mm; condyle clearance 43.1 mm at 10 deg of "
     "extension; moment arm spans 59.4 to 102.8 mm (mean 72.4 mm); margin vs plain "
     "human target +39.77% (optimizer winner before the final adjustment: "
     "+31.30% / +42.34%)."),
    ("Direct flexor",
     "-81.23 to -3.864", "-23.16 to -2.461", 5.01, -107,
     "Both actuators share one attachment route; the mirrored pair is the next "
     "design refinement; applies the adopted correction terms including the "
     "extensor wrapping factor."),
    ("Mono-pulley flexor variant",
     "peak -85.74 (insertion-side -85.56)", "-23.16 to -2.461", "not reported",
     "not reported",
     "Pulley gain fixed at 1:1, a redirect rather than a force multiplier; deepens "
     "the peak predicted flexion torque by 4.51 N\u00b7m (5.6%) over the direct "
     "route while holding the off-axis torque to at most 0.803 N\u00b7m; evaluates "
     "its torque requirement on the insertion-side torque; objective values of the "
     "direct and pulley formulations are not commensurable."),
]

# ----------------------------------------------------------------------
# T5 - Bi-pulley batch (chapter tab:bipulley_batch + ledger 6.5 clearances)
# ----------------------------------------------------------------------
T5_ROWS = [
    ("soleus", 0.9968,
     "Converged (pattern-search exit flag 1), nonpositive torque margin", 19.142),
    ("tibialis anterior", 0.9288,
     "Converged (pattern-search exit flag 1), nonpositive torque margin", 15.558),
    ("rectus femoris", 3.423,
     "Nonpositive torque margin (worst constraint 0.2423)", 7.024),
    ("medial gastrocnemius", 101,
     "Pattern-search exit flag -2, infeasibility sentinel (worst constraint 10)",
     "none recorded"),
    ("biceps femoris short head", 101,
     "Pattern-search exit flag -2, infeasibility sentinel (worst constraint 10)",
     4.260),
]


def main():
    before = set(os.listdir(OUT_DIR))

    gof_header = ["Test",
                  "Improved BPA char. alone: RMSE (" + NM + ")",
                  "Improved BPA char. alone: FVU",
                  "Improved BPA char. alone: Max. (" + NM + ")",
                  "Corrected model: RMSE (" + NM + ")",
                  "Corrected model: FVU",
                  "Corrected model: Max. (" + NM + ")"]

    cap1 = ("Pinned-knee flexor torque-prediction goodness of fit per test. "
            "Improved-BPA-characterization-alone = improved BPA force "
            "characterization with rigid attachments, no length offset, and no "
            "wrapping loss; corrected model = adopted two-stage correction terms. "
            "RMSE and maximum absolute residual (Max.) are in N\u00b7m; FVU = "
            "fraction of variance unexplained. The l0 = 47.9 cm test was held out "
            "of training; the training mean averages the four training tests.")
    cap2 = ("Pinned-knee extensor torque-prediction goodness of fit per test, at "
            "the adopted correction terms with the stiffness pair locked to the "
            "flexor-identified values. Column definitions as in the flexor table. "
            "RMSE and maximum absolute residual (Max.) are in N\u00b7m; FVU = "
            "fraction of variance unexplained. The extensor terms were fitted over "
            "the fitting tests, the l0 = 45.7 cm test was held out, and the fitting "
            "mean averages the five fitting tests.")
    cap3 = ("Adopted correction terms of the two-stage identification. The "
            "extensor refit locked (chi1, chi2) to the flexor-identified values, "
            "so one stiffness pair is shared across configurations. The pair was "
            "adopted from a family of near-equivalent fits by a requirement "
            "external to the identification: both route-redesign configurations "
            "must meet the required 5% torque-overshoot target simultaneously. "
            "Full-precision values are the underlying identified values; the "
            "reported values match the manuscript text.")
    cap4 = ("Route-redesign outcomes against the human isometric torque targets "
            "under the shared actuator plan of two 20 mm BPAs per route, a 620 kPa "
            "pressure ceiling, and a maximum usable contraction of 0.255. Negative "
            "torques denote knee flexion. The margin is stated with respect to the "
            "required envelope (human target + 5% overshoot). 'Not reported' marks "
            "a quantity not of record for that variant; for the mono-pulley variant "
            "only the peak torque is quoted rather than a span.")
    cap5 = ("Bi-pulley batch on five muscle sites of the robot-body model's "
            "OpenSim leg. Objectives are unitless; 101 is the batch's "
            "infeasibility sentinel; all five runs closed with a nonpositive "
            "torque margin. Adjacent routes are required to clear one another by "
            "25 mm (two 10 mm actuator radii + 5 mm bundle margin); the "
            "medial-gastrocnemius run records no finite minimum clearance.")

    written = []
    # of-record printed precision: RMSE 3 dp, FVU 4 dp, Max 3 dp (cols 2..7)
    gof_fmt = {c: ("0.000" if c in (2, 5) else "0.0000" if c in (3, 6) else "0.000")
               for c in range(2, 8)}
    written.append(build(
        os.path.join(OUT_DIR, "PinnedFlexor_GoF.xlsx"), cap1, gof_header,
        T1_ROWS, [30, 13, 12, 13, 12, 12, 12], mean_flags={5, 6},
        num_formats=gof_fmt))
    written.append(build(
        os.path.join(OUT_DIR, "PinnedExtensor_GoF.xlsx"), cap2, gof_header,
        T2_ROWS, [30, 13, 12, 13, 12, 12, 12], mean_flags={6, 7},
        num_formats=gof_fmt))
    written.append(build(
        os.path.join(OUT_DIR, "AdoptedCorrectionTerms.xlsx"), cap3,
        ["Parameter", "Configuration", "Role", "Value (full precision)", "Unit",
         "Value as reported", "Search bound, lower", "Search bound, upper",
         "Bounds unit"],
        T3_ROWS, [10, 13, 42, 18, 9, 14, 16, 16, 10]))
    written.append(build(
        os.path.join(OUT_DIR, "RouteRedesign_Summary.xlsx"), cap4,
        ["Route variant", "Predicted torque (" + NM + ")",
         "Human-scaled requirement (" + NM + ")", "Minimum margin (%)",
         "Worst angle (deg)", "Note"],
        T4_ROWS, [26, 24, 20, 14, 13, 60]))
    written.append(build(
        os.path.join(OUT_DIR, "BiPulley_Batch.xlsx"), cap5,
        ["Muscle site", "Objective (unitless)", "Outcome",
         "Minimum route clearance (mm)"],
        T5_ROWS, [26, 16, 56, 20], num_formats={4: "0.000"}))

    print("WROTE:")
    for p in written:
        print("  " + p + "  (" + str(os.path.getsize(p)) + " bytes)")

    # ---------------- verification 1: mean rows recompute ----------------
    tol = 5.1e-4  # printed cells carry <=5e-4 rounding each
    failures = []

    def check_means(name, rows, train_idx, mean_row):
        for col in range(1, 7):
            vals = [rows[i][col] for i in train_idx]
            got = sum(vals) / len(vals)
            want = rows[mean_row][col]
            ok = abs(got - want) <= tol
            print("  %s mean col%d: recomputed %.6f vs of-record %.4f -> %s"
                  % (name, col, got, want, "OK" if ok else "MISMATCH"))
            if not ok:
                failures.append((name, col, got, want))

    print("\nMEAN-ROW RECOMPUTATION (tolerance %.1e, printed-cell precision):" % tol)
    # T1: training rows = 48.5, 45.7, 40.6-tendon, 41.0 (indices 0,1,3,4); mean row 5
    check_means("T1-flexor", T1_ROWS, [0, 1, 3, 4], 5)
    # T2: fitting rows = 40.0, 40.0-tendon, 43.6, 43.2-tendon, 46.5 (0,1,2,3,5); mean row 6
    check_means("T2-extensor", T2_ROWS, [0, 1, 2, 3, 5], 6)
    # validation means must equal their single held-out rows
    for name, rows, val_row, mean_row in (("T1-flexor", T1_ROWS, 2, 6),
                                          ("T2-extensor", T2_ROWS, 4, 7)):
        for col in range(1, 7):
            if rows[val_row][col] != rows[mean_row][col]:
                failures.append((name + "-valmean", col,
                                 rows[mean_row][col], rows[val_row][col]))
        print("  %s validation mean equals its held-out row: %s"
              % (name, "OK" if not any(f[0].startswith(name + "-valmean")
                                       for f in failures) else "MISMATCH"))

    # ---------------- verification 2: readback ----------------
    print("\nREADBACK:")
    for p in written:
        wb = load_workbook(p)
        ws = wb.active
        print("  %s: sheet '%s', %d rows x %d cols (incl. caption+header)"
              % (os.path.basename(p), ws.title, ws.max_row, ws.max_column))
        for r in range(3, ws.max_row + 1):
            vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            print("    " + " | ".join(str(v) for v in vals))

    # ---------------- verification 3: directory diff ----------------
    after = set(os.listdir(OUT_DIR))
    new = sorted(after - before)
    print("\nDIRECTORY DIFF (new files in Tables/): " + str(new))
    if any(not f.lower().endswith(".xlsx") for f in new):
        failures.append(("non-xlsx created", new, None, None))

    if failures:
        print("\nFAILURES: " + repr(failures))
        sys.exit(1)
    print("\nALL CHECKS PASSED")


if __name__ == "__main__":
    main()
