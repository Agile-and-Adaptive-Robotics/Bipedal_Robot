from openpyxl import load_workbook

path = "Documentation/Reports and Papers/Knee_Torque_Test/Tables/PinnedFlexor_Preopt_GoF.xlsx"
repl = {
    (4.871, 0.8354, 9.865): (7.617, 2.1278, 25.531),
    (6.074, 2.7090, 9.262): (9.694, 2.1149, 21.076),
    (4.465, 0.6640, 9.536): (8.803, 1.9380, 26.264),
    (3.226, 0.4512, 7.376): (6.373, 0.5935, 16.859),
}
hits = 0
wb = load_workbook(path)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        vals = [c.value for c in row]
        for old, new in repl.items():
            if all(any(isinstance(v, (int, float)) and abs(v - o) < 5e-4 for v in vals) for o in old):
                idx = []
                for o in old:
                    for c in row:
                        if isinstance(c.value, (int, float)) and abs(c.value - o) < 5e-4:
                            idx.append(c)
                            break
                for c, n in zip(idx, new):
                    c.value = n
                hits += 1
wb.save(path)
print(f"updated rows: {hits} (expect 4)")
