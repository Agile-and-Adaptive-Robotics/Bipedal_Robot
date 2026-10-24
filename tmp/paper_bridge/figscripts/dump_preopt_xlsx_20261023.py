from openpyxl import load_workbook

wb = load_workbook("Documentation/Reports and Papers/Knee_Torque_Test/Tables/PinnedFlexor_Preopt_GoF.xlsx")
for sh in wb.worksheets:
    print(f"=== sheet: {sh.title} ({sh.max_row}x{sh.max_column}) ===")
    for r in sh.iter_rows():
        print(" ", [c.value for c in r])
