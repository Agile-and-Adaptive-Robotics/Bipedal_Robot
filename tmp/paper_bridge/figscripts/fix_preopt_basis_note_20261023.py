from openpyxl import load_workbook

path = "Documentation/Reports and Papers/Knee_Torque_Test/Tables/PinnedFlexor_Preopt_GoF.xlsx"
new_note = ("Basis: each prediction is scored against all measured points at which it is "
            "defined; the Hunt-equation prediction is undefined near full extension, where "
            "its modeled strain falls below zero. Torque residuals in N*m. Matches "
            "tab:FlxPinPreGof, Main.tex.")
wb = load_workbook(path)
sh = wb["Preopt GoF"]
sh.cell(row=11, column=1).value = new_note
wb.save(path)
print("basis note updated")
