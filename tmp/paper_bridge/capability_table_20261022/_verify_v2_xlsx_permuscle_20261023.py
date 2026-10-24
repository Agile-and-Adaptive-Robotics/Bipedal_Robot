# _verify_v2_xlsx_permuscle_20261023.py -- xlsx 'Per muscle' sheet cell-by-cell vs
# per_muscle_capability_20261023.csv; 'Capability 2x2' sheet re-checked strictly.
import openpyxl, csv, os

BASE = os.path.dirname(os.path.abspath(__file__))
problems = []
def bad(m): problems.append(m)

wb = openpyxl.load_workbook(os.path.join(BASE, "capability_table_20261023.xlsx"), data_only=True)

# ---- 'Capability 2x2' strict check ----
ws = wb["Capability 2x2"]
want = {
    "A1": "Row", "B1": "% of RoM torque meets or exceeds human", "C1": "% of RoM with active force",
    "A2": "Hand-tuned (27 of 27 graded)", "B2": 72.49, "C2": 96.64,
    "A3": "Optimized (27 of 27 graded; 27 built)", "B3": 65.19, "C3": 74.27,
}
for coord, val in want.items():
    got = ws[coord].value
    ok = abs(got - val) < 1e-9 if isinstance(val, float) else got == val
    if not ok:
        bad(f"Capability 2x2!{coord}: xlsx {got!r} vs expected {val!r}")
# no other numeric cells on this sheet beyond B2 C2 B3 C3
for row in ws.iter_rows():
    for c in row:
        if isinstance(c.value, (int, float)) and c.coordinate not in ("B2", "C2", "B3", "C3"):
            bad(f"unexpected numeric on Capability 2x2: {c.coordinate} = {c.value!r}")
print("Capability 2x2 sheet: strict cell check done")

# ---- 'Per muscle' sheet vs CSV ----
ws2 = wb["Per muscle"]
with open(os.path.join(BASE, "per_muscle_capability_20261023.csv"), newline="", encoding="utf-8") as f:
    rows = list(csv.reader(f))
header, data = rows[0], rows[1:]
xl_header = [c.value for c in ws2[1]]
if xl_header != header:
    bad(f"Per-muscle header mismatch: xlsx {xl_header} vs csv {header}")
if ws2.max_row - 1 != len(data):
    bad(f"Per-muscle row count: xlsx {ws2.max_row - 1} vs csv {len(data)}")

numeric_cols = {i for i, h in enumerate(header) if any(
    s in h for s in ("pct", "cells", "flags", "nBPA", "_G", "size_mm", "evals", "actuator_id", "converged"))}
mismatch = 0
for r, drow in enumerate(data, start=2):
    for ci, csv_val in enumerate(drow):
        xl_val = ws2.cell(row=r, column=ci + 1).value
        if ci in numeric_cols:
            try:
                ok = abs(float(xl_val) - float(csv_val)) < 1e-9
            except (TypeError, ValueError):
                ok = str(xl_val) == csv_val
        else:
            ok = str(xl_val) == csv_val
        if not ok:
            mismatch += 1
            if mismatch <= 20:
                bad(f"Per muscle r{r} {header[ci]}: xlsx {xl_val!r} vs csv {csv_val!r}")
print(f"Per-muscle sheet: {len(data)} rows x {len(header)} cols compared, {mismatch} mismatches")

print("\n===== NO PROBLEMS =====" if not problems else "\n===== PROBLEMS =====")
for m in problems: print("[DISCREPANCY]", m)
