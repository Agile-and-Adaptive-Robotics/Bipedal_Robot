"""Read summary tabs of the 2026_06_Festo results tables (READ-ONLY)."""
import openpyxl

BASE = r"D:\GitHub\Bipedal_Robot\Testing_Data\2026_06_Festo"

for fname in ["Results_table_10mm_pinned.xlsx", "Results_table_20mm.xlsx"]:
    path = BASE + "\\" + fname
    print("\n############ " + fname + " ############")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    print("sheets:", wb.sheetnames)
    for ws in wb.worksheets:
        print(f"\n===== sheet '{ws.title}' ({ws.max_row}x{ws.max_column}) =====")
        nrows = min(ws.max_row, 40)
        for r, row in enumerate(ws.iter_rows(min_row=1, max_row=nrows, values_only=True), start=1):
            cells = ["" if c is None else (round(c, 4) if isinstance(c, float) else c) for c in row]
            # trim trailing empties
            while cells and cells[-1] == "":
                cells.pop()
            if not cells:
                continue
            print(f"  r{r}: " + " | ".join(str(c) for c in cells))
        if ws.max_row > nrows:
            print(f"  ... ({ws.max_row - nrows} more rows)")
    wb.close()
