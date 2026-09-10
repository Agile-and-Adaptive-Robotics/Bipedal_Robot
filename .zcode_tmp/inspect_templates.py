import sys
from pptx import Presentation
from pptx.util import Emu

def inches(v):
    return round(Emu(v).inches, 2) if v is not None else None

def fontinfo(r):
    f = r.font
    parts = []
    if f.name: parts.append(f.name)
    if f.size: parts.append(f"{f.size.pt:.0f}pt")
    if f.bold: parts.append("B")
    if f.italic: parts.append("I")
    try:
        if f.color and f.color.type is not None:
            if str(f.color.type) == "MSO_THEME_COLOR.NOT_THEME_COLOR" or "RGB" in str(f.color.type):
                parts.append(f"#{f.color.rgb}")
            else:
                parts.append(f"theme:{f.color.theme_color}")
    except Exception:
        pass
    return ",".join(parts)

def walk(shapes, depth=0):
    for s in shapes:
        ind = "  " * depth
        st = str(s.shape_type).split(".")[-1].split(" ")[0] if s.shape_type is not None else "?"
        ph = ""
        if s.is_placeholder:
            ph = f" [PH {s.placeholder_format.type} idx={s.placeholder_format.idx}]"
        print(f"{ind}- {st} '{s.name}'{ph} x={inches(s.left)} y={inches(s.top)} w={inches(s.width)} h={inches(s.height)}")
        if s.shape_type == 6:  # group
            walk(s.shapes, depth + 1)
            continue
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                t = p.text.strip()
                if not t: continue
                fonts = set()
                for r in p.runs:
                    fonts.add(fontinfo(r))
                print(f"{ind}    TXT: {t[:80]!r} | {sorted(fonts)}")
        if st == "PICTURE":
            try:
                print(f"{ind}    IMG: {s.image.filename or s.image.content_type} {s.image.size}")
            except Exception:
                pass

for path in sys.argv[1:]:
    print("=" * 100)
    print("FILE:", path)
    prs = Presentation(path)
    W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    print(f"SLIDE SIZE: {W:.2f} x {H:.2f} in  ({len(prs.slides)} slides, {len(prs.slide_masters)} masters, {len(prs.slide_layouts)} layouts)")
    for i, layout in enumerate(prs.slide_layouts):
        phs = [f"{p.placeholder_format.type}(idx={p.placeholder_format.idx})" for p in layout.placeholders]
        print(f"  layout[{i}] '{layout.name}': {phs if phs else 'no placeholders'}")
    for idx, slide in enumerate(prs.slides):
        print("-" * 90)
        print(f"SLIDE {idx} (layout='{slide.slide_layout.name}')")
        walk(slide.shapes)
