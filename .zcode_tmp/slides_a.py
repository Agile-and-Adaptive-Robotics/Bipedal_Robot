"""Slides 1-14: title, framing, Aim 1 (BPA force characterization)."""
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from deck_lib import *

LOGO = os.path.join(ASSETS, "aarl_logo_0.png")

def s01_title(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    pic(s, LOGO, box=(4.47, 0.75, 4.4, 1.34))
    tx(s, 0.7, 2.55, 11.93, 1.9, "Movement and Control of\nBiomimetic Humanoid Robots",
       size=40, color=BROWN, bold=True, align=PP_ALIGN.CENTER, spacing=1.0)
    tx(s, 0.7, 4.45, 11.93, 0.45, "Dissertation Defense", size=20, color=MUTED, align=PP_ALIGN.CENTER)
    tx(s, 0.7, 5.15, 11.93, 0.5, "Ben P. Bolen", size=24, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    tx(s, 0.7, 5.72, 11.93, 0.4, "Advisor: M. Anthony Hunt, Ph.D.", size=15, color=MUTED, align=PP_ALIGN.CENTER)
    tx(s, 0.7, 6.30, 11.93, 0.4, "Agile & Adaptive Robotics Laboratory  ·  Portland State University  ·  September 2026",
       size=13, color=MUTED, align=PP_ALIGN.CENTER)
    s.notes_slide.notes_text_frame.text = (
        "Good morning. Committee members, thank you for reading the dissertation. "
        "Today: can pneumatic artificial muscles be characterized and arranged so a biomimetic robot "
        "reproduces human joint force capability - and how the controller will be built on top.")

def s02_roadmap(prs, n):
    s = content_slide(prs, "Roadmap",
        "Three questions connect one actuator to a walking, neurally controlled robot",
        cite="Dissertation Ch. 3-6; Bolen et al. 2026, Actuators; Bolen et al., in preparation 2025",
        notes="One sentence per aim. Aim 1 published; Aim 2 in manuscript form; Aim 3 is the defined path.")
    items = [
        ("01", "What force can one\nactuator produce?", "Isometric characterization of Festo BPAs:\nresting length, pressure, contraction", "537 tests  ·  dimensionless force surface", GREEN),
        ("02", "What torque can a\nBPA-driven knee produce?", "Pinned + biomimetic knees; correction\nterms identified by multiobjective fit", "meets / exceeds human torque over\nmost of range of motion", GREEN),
        ("03", "How will the robot\nbe controlled?", "OpenSim to MuJoCo pipeline; two-layer\nspinal CPG with full proprioception", "rhythm verified  ·  closed loop next", TEAL),
    ]
    x = MX; w = 3.85; gap = 0.35
    for i, (num, q, m, tag, c) in enumerate(items):
        cx = x + i * (w + gap)
        card(s, cx, 2.0, w, 4.55, fill=CARD if i < 2 else CARD2)
        tx(s, cx + 0.3, 2.25, 1.6, 0.9, num, size=40, color=BROWN, bold=True)
        tx(s, cx + 0.3, 3.15, w - 0.6, 1.0, q, size=17, color=TEXT, bold=True, spacing=1.0)
        tx(s, cx + 0.3, 4.35, w - 0.6, 1.1, m, size=13, color=MUTED, spacing=1.05)
        tx(s, cx + 0.3, 5.75, w - 0.6, 0.7, tag, size=13, color=c, bold=True, spacing=1.0)
        if i < 2:
            arrow(s, cx + w + 0.03, 4.25, cx + w + gap - 0.03, 4.25, color=FAINT, width=2.2)
    pageno(s, n)

def s03_project(prs, n):
    s = content_slide(prs, "The project",
        "A humanoid robot from the trunk down: human-like body plan, synthetic nervous system",
        cite="Bolen & Hunt 2019; Morrow et al. 2020; Steele 2018; Scharzenberger et al. 2019",
        notes="Ground the committee: this dissertation is the actuator-and-joint foundation of the AARL humanoid program.")
    pic(s, os.path.join(FIG, "mujoco_robot.png"), box=(0.45, 1.75, 3.4, 5.1))
    tx(s, 0.45, 6.55, 3.5, 0.4, "converted leg model, 92 muscle paths", size=11, color=FAINT, italic=True)
    rows = [
        ("Guiding hypothesis", "human-like muscle numbers, architecture, routing + human-like torque capability, controlled by a biologically grounded neural architecture  \u2192  biomimetic movement", None),
        ("Built on", "my master's work: Gait2392-matched kinematics, muscle routing tools  \u00b7  Morrow: optimized attachments  \u00b7  Steele: biomimetic 4-bar knee  \u00b7  Scharzenberger: BPA quadruped", None),
        ("This dissertation", "the quantitative basis: characterize the actuator, correct the joint model against experiment, validate against human torque, define the control path", GREEN),
    ]
    y = 1.95
    for t, b, c in rows:
        rich(s, 4.35, y, 5.45, 1.5, [
            {"t": t.upper(), "size": 12, "color": TEAL if not c else c, "bold": True, "space_after": 3},
            {"t": b, "size": 14.5, "color": TEXT, "spacing": 1.08},
        ])
        y += 1.62
    pic(s, os.path.join(MEDIA, "image44.jpg"), box=(10.0, 5.45, 2.8, 1.5))
    tx(s, 4.35, 6.55, 5.45, 0.4, "AARL quadruped: BPA actuators + SNS, already on a treadmill", size=11, color=FAINT, italic=True)
    pageno(s, n)

def s04_motivation(prs, n):
    s = content_slide(prs, "Why biomimetic robots",
        "Robots run neuromechanical experiments that are impossible in living subjects",
        cite="Ijspeert 2020; Shin et al. 2018; photo: Asano et al. 2016, Humanoids",
        notes="The biology professor on the committee will recognize this loop: model, perturb, understand.")
    node(s, MX, 2.15, 3.5, 1.15, "Biology\nanatomy \u00b7 architecture \u00b7 afferents", fill=CARD2, size=14, bold=False)
    node(s, MX, 4.9, 3.5, 1.15, "Robot\nactuators \u00b7 routing \u00b7 controller", fill=CARD, size=14)
    arrow(s, MX + 1.0, 3.4, MX + 1.0, 4.8, color=GREEN, width=2.5)
    tx(s, MX + 1.15, 3.85, 2.6, 0.5, "design principles", size=12, color=GREEN, bold=True)
    arrow(s, MX + 2.6, 4.8, MX + 2.6, 3.4, color=BROWN, width=2.5)
    tx(s, MX + 2.72, 3.85, 2.6, 0.5, "test hypotheses", size=12, color=BROWN, bold=True)
    pic(s, os.path.join(MEDIA, "image11.png"), box=(8.9, 1.9, 3.9, 4.8))
    tx(s, 8.9, 6.72, 3.9, 0.35, "human-mimetic musculoskeletal humanoid", size=11, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    exps = [
        "Remove one muscle from the system",
        "Block one sensory feedback pathway",
        "Stimulate a named neural pathway",
        "Identical copies - no individual variance",
        "Faster, cheaper, ethical vs. human studies",
    ]
    y = 2.1
    for e in exps:
        tx(s, 4.55, y, 4.2, 0.5, "\u2013  " + e, size=16, color=TEXT, spacing=1.0)
        y += 0.62
    tx(s, 4.55, 5.55, 4.2, 1.2, "Loss-of-function and stimulation experiments\nare the point of the platform.", size=15, color=BROWN, bold=True, spacing=1.05)
    pageno(s, n)

def s05_bpa(prs, n):
    s = content_slide(prs, "The actuator",
        "Braided pneumatic actuators: muscle-like strengths, and three modeling headaches",
        cite="Hunt et al. 2017; S\u00e1rosi et al. 2012-2017; Martens & Boblan 2017; Festo DMSP datasheets",
        notes="Sets up gap 1. Force peaks AT resting length - unlike a biological muscle whose optimum is mid-range.")
    pic(s, os.path.join(MEDIA, "image5.png"), box=(0.55, 1.95, 2.1, 4.6))
    tx(s, 0.35, 6.6, 2.5, 0.35, "Festo DMSP-10/20", size=11, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    plus = [
        "High force-to-weight, low mass",
        "Intrinsic compliance - safe, spring-like",
        "Tension-only, like biological muscle",
        "Force-length curve qualitatively muscle-like",
    ]
    minus = [
        "Maximum force occurs AT resting length",
        "Contraction limited (max \u2248 \u03b5\u2086\u2082\u2080)",
        "F highly nonlinear in l\u2080, P, \u03b5",
        "Manufacturer tool mis-predicts short BPAs",
    ]
    rich(s, 3.3, 1.95, 4.6, 0.4, [{"t": "WHY WE USE THEM", "size": 13, "color": GREEN, "bold": True}])
    y = 2.45
    for p_ in plus:
        tx(s, 3.3, y, 4.6, 0.5, "+  " + p_, size=15.5, color=TEXT, spacing=1.0); y += 0.68
    rich(s, 8.25, 1.95, 4.5, 0.4, [{"t": "WHY THEY ARE HARD", "size": 13, "color": RED, "bold": True}])
    y = 2.45
    for m_ in minus:
        tx(s, 8.25, y, 4.5, 0.5, "\u2013  " + m_, size=15.5, color=TEXT, spacing=1.0); y += 0.68
    card(s, 3.3, 5.35, 9.45, 1.35, fill=CARD)
    rich(s, 3.6, 5.55, 8.9, 1.0, [
        {"t": "Designing a biomimetic robot around BPAs needs quantitative, predictive models:", "size": 14.5, "color": MUTED, "space_after": 3},
        {"t": "actuator force  \u2192  joint torque  \u2192  comparison with human capability", "size": 16, "color": BROWN, "bold": True},
    ])
    pageno(s, n)

def s06_gaps(prs, n):
    s = content_slide(prs, "Problem statement",
        "Three gaps stood between the existing design and a robot that walks",
        cite="Dissertation Ch. 1 \u00a71.3",
        notes="Each gap maps to one aim. The committee heard these as the research objectives.")
    rows = [
        ("GAP 1", "Actuator force models inadequate for design",
         "Festo tool over-predicts short BPAs; published models fit a single length each",
         GREEN), 
        ("GAP 2", "Actuator force \u2260 joint torque",
         "brackets, artificial tendons, wrap paths: unknown losses; optimized biomimetic routing never validated vs. human torque",
         GREEN),
        ("GAP 3", "No pipeline from musculoskeletal model to biologically realistic controller",
         "two-level CPG ingredients existed, never assembled for a BPA biped with human-like routing",
         TEAL),
    ]
    y = 1.9
    for tag, head, sub, c in rows:
        card(s, MX, y, 12.2, 1.38, fill=CARD)
        tx(s, MX + 0.28, y + 0.2, 1.2, 0.5, tag, size=15, color=c, bold=True)
        rich(s, MX + 1.55, y + 0.16, 10.4, 1.1, [
            {"t": head, "size": 17, "color": TEXT, "bold": True, "space_after": 2},
            {"t": sub, "size": 13.5, "color": MUTED, "spacing": 1.0},
        ])
        y += 1.52
    tx(s, MX, 6.5, 12.2, 0.4, "Aims 1-3 mirror the gaps  \u2192", size=14, color=BROWN, bold=True)
    pageno(s, n)

def s07_jig(prs, n):
    s = content_slide(prs, "Aim 1 \u00b7 actuator characterization",
        "A custom jig measured isometric force across resting length, pressure, and contraction",
        cite="Bolen, Elzein, Pang, Hunt 2026, Actuators (Ch. 3)",
        notes="Hardware slide - one breath. The point: wide l0 range, both diameters, 537 pairs.")
    pic(s, os.path.join(FIG, "testJigs_force.png"), box=(0.7, 1.9, 4.4, 4.75))
    specs = [
        ("Diameters", "10 mm and 20 mm Festo DMSP"),
        ("Resting lengths", "112-518 mm (10 mm)  \u00b7  300-509 mm (20 mm)"),
        ("Pressure", "up to P\u2086\u2082\u2080 = 620 kPa supply"),
        ("Data", "537 force-pressure-contraction pairs (321 + 216)"),
        ("Instrumentation", "S-type load cells + HX711 \u00b7 MPX5700 pressure \u00b7 Arduino \u2192 MATLAB"),
    ]
    y = 2.1
    for k, v in specs:
        rich(s, 5.6, y, 7.1, 0.9, [
            {"t": k.upper(), "size": 11.5, "color": TEAL, "bold": True, "space_after": 1},
            {"t": v, "size": 15.5, "color": TEXT, "spacing": 1.0},
        ])
        y += 0.95
    pageno(s, n)

def s08_fmax_discovery(prs, n):
    s = content_slide(prs, "Aim 1 \u00b7 actuator characterization",
        "Maximum isometric force depends on resting length - a previously unreported characteristic",
        cite="Bolen et al. 2026, Actuators, Figs. 2-3",
        notes="The discovery slide. Festo's own tool predicts the OPPOSITE trend - force rising as length shrinks.")
    pic(s, os.path.join(FIG, "MaxForce10.png"), box=(0.7, 1.9, 5.9, 3.95))
    pic(s, os.path.join(FIG, "MaxForce20.png"), box=(6.75, 1.9, 5.9, 3.95))
    tx(s, 0.7, 5.92, 5.9, 0.4, "\u03c6 10 mm:  F\u2086\u2082\u2080 vs l\u2080", size=13, color=MUTED, align=PP_ALIGN.CENTER)
    tx(s, 6.75, 5.92, 5.9, 0.4, "\u03c6 20 mm:  F\u2086\u2082\u2080 vs l\u2080", size=13, color=MUTED, align=PP_ALIGN.CENTER)
    card(s, 0.7, 6.38, 12.0, 0.52, fill=CARD2)
    tx(s, 1.0, 6.48, 11.5, 0.4,
       "solid: our arctan fit (data-driven)  \u00b7  dashed: Festo tool - predicts force INCREASING as l\u2080 shrinks",
       size=13.5, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    pageno(s, n)

def s09_fmax_model(prs, n):
    s = content_slide(prs, "Aim 1 \u00b7 actuator characterization",
        "Two coefficients per diameter capture the resting-length effect",
        cite="Bolen et al. 2026, Actuators, Eqs. 4-6; fit: nonlinear least squares, least-absolute-residual",
        notes="Math committee: arctan form, bounded asymptote, physical offset from end-cap contact in CAD.")
    card(s, MX, 1.95, 7.0, 1.5, fill=CARD)
    eq(s, "fmax", box=(1.0, 2.12, 6.3, 1.16))
    card(s, MX, 3.6, 7.0, 1.5, fill=CARD)
    eq(s, "fmax20", box=(1.0, 3.77, 6.3, 1.16))
    card(s, MX, 5.25, 7.0, 1.45, fill=CARD2)
    eq(s, "fmax3d", box=(1.0, 5.45, 6.3, 1.05))
    rich(s, 8.0, 2.0, 4.75, 4.8, [
        {"t": "LIMITS BEHAVE", "size": 12, "color": TEAL, "bold": True, "space_after": 5},
        {"t": "l\u2080 \u2192 \u221e:   470 N (10 mm) vs Festo 490 N", "size": 15.5, "color": TEXT, "space_after": 1},
        {"t": "within Festo's stated 10% manufacturing tolerance (20 mm: 1460 vs 1570 N)", "size": 12.5, "color": MUTED, "space_after": 10},
        {"t": "THE OFFSET IS PHYSICAL", "size": 12, "color": TEAL, "bold": True, "space_after": 5},
        {"t": "7.5 mm / 13 mm = length at which the end caps contact each other in the CAD model - no contractile stroke remains", "size": 14.5, "color": TEXT, "spacing": 1.1, "space_after": 10},
        {"t": "As l\u2080 \u2192 0, measured force \u2192 0.", "size": 15.5, "color": BROWN, "bold": True},
    ])
    pageno(s, n)

def s10_festo_wrong(prs, n):
    s = content_slide(prs, "Aim 1 \u00b7 actuator characterization",
        "At short resting lengths the manufacturer's tool mis-predicts force by 42%",
        cite="Bolen et al. 2026, Actuators, Discussion \u00a71; example: \u03c610 mm, l\u2080 = 112 mm, 620 kPa",
        notes="The design-relevant failure: anyone specifying BPAs under 300 mm gets burned by the tool.")
    bar_chart(s, 0.65, 2.0, 6.3, 4.6,
              ["Measured", "Our model", "Festo tool"], [350.9, 335.3, 498.6],
              [TEAL, GREEN, RED], title="\u03c610 mm at l\u2080 = 112 mm, 620 kPa:  isometric force (N)",
              max_scale=600)
    tx(s, 0.65, 6.65, 6.3, 0.4, "error:  4.5% (ours)  vs  42.1% (Festo)", size=15, color=TEXT,
       bold=True, align=PP_ALIGN.CENTER)
    card(s, 7.35, 2.1, 5.35, 4.3, fill=CARD)
    rich(s, 7.65, 2.35, 4.75, 3.9, [
        {"t": "WHY THE TOOL FAILS SHORT", "size": 12, "color": TEAL, "bold": True, "space_after": 6},
        {"t": "Ideal cylindrical models ignore end effects; the clamped ends force steep stress gradients.", "size": 14.5, "color": TEXT, "spacing": 1.12, "space_after": 8},
        {"t": "The end-effect zone has roughly fixed length - so as l\u2080 shrinks, it dominates.", "size": 14.5, "color": TEXT, "spacing": 1.12, "space_after": 8},
        {"t": "Measured force \u2192 0 as l\u2080 \u2192 0.  The Festo tool instead predicts exponential growth.", "size": 14.5, "color": RED, "spacing": 1.12, "space_after": 8},
        {"t": "\u201cResearchers using BPA resting lengths under 300 mm should take note.\u201d", "size": 14, "color": BROWN, "bold": True, "italic": True, "spacing": 1.1},
    ])
    pageno(s, n)

def s11_surface(prs, n):
    s = content_slide(prs, "Aim 1 \u00b7 actuator characterization",
        "All 537 measurements collapse onto one dimensionless force surface",
        cite="Bolen et al. 2026, Actuators, Eqs. 7-9; \u03b5* = \u03b5/\u03b5\u2086\u2082\u2080,  P* = P/620 kPa",
        notes="Normalization by each muscle's own F620 and eps620 is what makes one surface fit all lengths.")
    card(s, MX, 1.9, 12.2, 1.05, fill=CARD2)
    eq(s, "fstar", 1.0, 2.12, h=0.62)
    pic(s, os.path.join(FIG, "FStar10.png"), box=(0.75, 3.15, 5.5, 3.55))
    pic(s, os.path.join(FIG, "FStar20.png"), box=(6.6, 3.15, 5.5, 3.55))
    tx(s, 0.75, 6.75, 5.5, 0.35, "F*(\u03b5*, P*) - \u03c610 mm, pressure isoclines", size=12, color=MUTED, align=PP_ALIGN.CENTER)
    tx(s, 6.6, 6.75, 5.5, 0.35, "F*(\u03b5*, P*) - \u03c620 mm", size=12, color=MUTED, align=PP_ALIGN.CENTER)
    pageno(s, n)

def s12_surface_gof(prs, n):
    s = content_slide(prs, "Aim 1 \u00b7 actuator characterization",
        "Three coefficients per diameter: force predicted out-of-sample to a few percent",
        cite="Bolen et al. 2026, Actuators, Table 2; 80/20 train/validation split",
        notes="Stress: validation is 20% held out, and max error under ~11% everywhere - at ANY length.")
    cols = [
        ("\u03c6 10 mm", "0.9998", "0.9994", "10.3%", "10.6%"),
        ("\u03c6 20 mm", "0.992", "0.9943", "6.8%", "5.7%"),
    ]
    hdr = ["", "adj. R\u00b2 (fit)", "adj. R\u00b2 (validation)", "max err (fit)", "max err (val.)"]
    x = MX; cw = [1.8, 2.6, 2.9, 2.4, 2.5]
    xx = x
    for i, h_ in enumerate(hdr):
        tx(s, xx + 0.1, 2.05, cw[i] - 0.2, 0.6, h_, size=13.5, color=TEAL, bold=True,
           align=PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER)
        xx += cw[i]
    y = 2.75
    for name, f, v, me, mv in cols:
        xx = x
        vals = [name, f, v, me, mv]
        for i, val in enumerate(vals):
            bold = i == 0 or i == 2
            colr = TEXT if i else BROWN
            tx(s, xx + 0.1, y, cw[i] - 0.2, 0.5, val, size=17 if i == 0 else 16, color=colr,
               bold=bold, align=PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER)
            xx += cw[i]
        y += 0.62
    tx(s, MX, 4.15, 12.2, 0.4, "RMSE (normalized):  0.005 (10 mm)  \u00b7  0.023 (20 mm)", size=14, color=MUTED)
    card(s, MX, 4.8, 12.2, 1.7, fill=CARD)
    rich(s, MX + 0.3, 5.0, 11.6, 1.4, [
        {"t": "The whole actuator model a designer needs:", "size": 15, "color": MUTED, "space_after": 4},
        {"t": "F*  (3 coefficients, per diameter)   \u00d7   F\u2086\u2082\u2080(l\u2080)  (2 coefficients, per diameter)", "size": 19, "color": BROWN, "bold": True},
    ])
    pageno(s, n)

def s13_vs_models(prs, n):
    s = content_slide(prs, "Aim 1 \u00b7 actuator characterization",
        "Lowest error of the compared models at every diameter and length tested",
        cite="Bolen et al. 2026, Actuators, Fig. 5 + Table 3; full table in Appendix C",
        notes="Sarosi params come from specific lengths - extrapolation is where they fail. Martens is geometry-sensitive.")
    pic(s, os.path.join(FIG, "Comparison.png"), box=(0.55, 1.9, 5.6, 4.95))
    tx(s, 0.55, 6.85, 5.6, 0.35, "measured vs predicted; y = x is perfect fit", size=11.5, color=FAINT,
       italic=True, align=PP_ALIGN.CENTER)
    rows = [
        ("10 mm, 257 mm", "22.9", "143.8", "73.8"),
        ("10 mm, 233 mm", "16.8", "132.7", "62.6"),
        ("20 mm, 300 mm", "35.6", "132.1", "72.8"),
        ("20 mm, 450 mm", "67.6", "177.7", "862.1"),
    ]
    tx(s, 6.6, 1.95, 6.2, 0.4, "RMSE (N) at matched conditions", size=13.5, color=TEAL, bold=True)
    cw = [2.1, 1.35, 1.35, 1.7]
    hdr = ["l\u2080", "Bolen", "S\u00e1rosi", "M&B"]
    xx = 6.6
    for i, h_ in enumerate(hdr):
        tx(s, xx, 2.4, cw[i], 0.4, h_, size=13.5, color=MUTED, bold=True,
           align=PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER)
        xx += cw[i]
    y = 2.85
    for r in rows:
        xx = 6.6
        for i, val in enumerate(r):
            if i == 1:
                tx(s, xx, y, cw[i], 0.45, val, size=16, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
            elif i == 0:
                tx(s, xx, y + 0.03, cw[i], 0.45, val, size=14, color=TEXT)
            else:
                tx(s, xx, y, cw[i], 0.45, val, size=15, color=MUTED, align=PP_ALIGN.CENTER)
            xx += cw[i]
        y += 0.58
    card(s, 6.6, 5.35, 6.2, 1.35, fill=CARD2)
    rich(s, 6.9, 5.52, 5.6, 1.1, [
        {"t": "Normalization travels: one fit per diameter works across lengths,", "size": 14, "color": TEXT, "spacing": 1.1},
        {"t": "while fixed-parameter models fail when moved off their fit length.", "size": 14, "color": TEXT, "spacing": 1.1},
    ])
    pageno(s, n)

def s14_realtime(prs, n):
    s = content_slide(prs, "Aim 1 \u2192 Aim 2",
        "The force model is two cheap maps - built to sit inside a real-time controller",
        cite="Bolen et al. 2026, Actuators, Discussion \u00a74; dynamics complement: Elzein et al. 2025",
        notes="Static map = feedforward plant model / lookup table. Valve-airflow dynamics are the Elzein layer.")
    card(s, 2.1, 2.1, 4.1, 1.5, fill=CARD2)
    tx(s, 2.3, 2.28, 3.7, 0.4, "MAP 1", size=12, color=TEAL, bold=True)
    tx(s, 2.3, 2.68, 3.7, 0.8, "F\u2086\u2082\u2080(l\u2080)\nresting length \u2192 max force", size=15.5, color=TEXT, spacing=1.05)
    card(s, 7.1, 2.1, 4.1, 1.5, fill=CARD2)
    tx(s, 7.3, 2.28, 3.7, 0.4, "MAP 2", size=12, color=TEAL, bold=True)
    tx(s, 7.3, 2.68, 3.7, 0.8, "F*(\u03b5*, P*)\nstrain, pressure \u2192 fraction", size=15.5, color=TEXT, spacing=1.05)
    arrow(s, 6.28, 2.85, 7.02, 2.85, color=FAINT)
    eq(s, "force", 3.6, 4.05, h=0.62)
    card(s, 2.1, 5.15, 9.1, 1.35, fill=CARD)
    rich(s, 2.4, 5.32, 8.6, 1.05, [
        {"t": "Closed form or lookup table; evaluated per timestep in the plant model.", "size": 15, "color": TEXT, "space_after": 3},
        {"t": "Next: multiply by geometry (moment arm) \u2192 joint torque.  Does it survive contact with the joint?", "size": 15, "color": BROWN, "bold": True},
    ])
    pageno(s, n)

def build(prs, start=2):
    s01_title(prs, 1)
    s02_roadmap(prs, 2)
    s03_project(prs, 3)
    s04_motivation(prs, 4)
    s05_bpa(prs, 5)
    s06_gaps(prs, 6)
    s07_jig(prs, 7)
    s08_fmax_discovery(prs, 8)
    s09_fmax_model(prs, 9)
    s10_festo_wrong(prs, 10)
    s11_surface(prs, 11)
    s12_surface_gof(prs, 12)
    s13_vs_models(prs, 13)
    s14_realtime(prs, 14)
    return 14
