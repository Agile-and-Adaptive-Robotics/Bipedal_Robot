"""Dump per-slide metadata (kicker, assertion, notes presence) for the QC gates."""
import json, os, sys
from pptx import Presentation

PPTX = r"D:\GitHub\Bipedal_Robot\Documentation\Reports and Papers\Dissertation\Defense\Bolen_Dissertation_Defense.pptx"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "slide_index.json")

prs = Presentation(PPTX)
rows = []
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip().replace("\n", " | ")[:110])
    notes = ""
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
    rows.append({"n": i, "texts": texts[:4], "notes": bool(notes),
                 "notes_head": notes[:140]})
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(rows, f, indent=1, ensure_ascii=False)
no_notes = [r["n"] for r in rows if not r["notes"]]
print(json.dumps({"slides": len(rows), "no_notes": no_notes}))
