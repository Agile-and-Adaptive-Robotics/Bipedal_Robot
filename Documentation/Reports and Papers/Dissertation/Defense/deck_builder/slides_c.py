"""Slides 26-41: Ch. 5 database, Ch. 6 SNS simulation, Ch. 7/8 close.

Assertion-evidence style per CONTENT_SPEC.md: chapter kickers (no Aim numbering),
one message per slide, figures carry the weight, speaker notes everywhere,
numbers anchored to ProofFinal chapter lines (see evidence_c.md).
"""
import os
from pptx.enum.text import PP_ALIGN
from deck_lib import *

LOGO = os.path.join(ASSETS, "aarl_logo_0.png")

# ---------------------------------------------------------------- Ch. 5

def s26_sadb_corpus(prs, n):
    s = content_slide(prs, "Ch. 5 \u00b7 Sensory Afferent Database",
        "A century of scattered afferent physiology, consolidated into one audited, offline-searchable resource",
        cite="Dissertation Ch. 5 (\u00a7A Structured Afferent Resource, \u00a7Audited Corpus and Curation Coverage); Bolen et al., SfN 2023",
        notes="Provenance: the database's three homes did not agree, so the 2026 campaign opened with a record-by-record "
              "reconciliation on DOI and normalized title (Ch. 5, line 14). No record was destroyed in the prune: every deletion "
              "was snapshotted, and the packaged explorer carries the 382 cut papers under an archived flag (line 20). "
              "Anticipated question: why build your own explorer app? Answer: an off-the-shelf bibliometric viewer was tried and "
              "retired within weeks - one crowded picture, no drill-through from cluster to papers, mega-cited classics dominating "
              "every view; this corpus needs filter-by-afferent-class, pathway-to-papers walks, and two-dimensional pivots (line 22).")
    pic(s, os.path.join(FIG, "c_sadb_explorer.png"), box=(0.55, 1.88, 6.6, 4.05))
    tx(s, 0.55, 6.02, 6.6, 0.7, "the packaged SADb Explorer: single file, runs offline in a browser \u00b7 table view with curation-state tags",
       size=11, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    stat(s, 7.55, 1.90, 2.5, "943", "papers in the audited corpus", vsize=34)
    stat(s, 10.25, 1.90, 2.5, "561", "retained papers (382 archived, none destroyed)", vsize=34)
    card(s, 7.55, 3.55, 5.2, 1.02, fill=CARD)
    rich(s, 7.75, 3.66, 4.85, 0.85, [
        {"t": "CITATION GRAPH", "size": 11, "color": TEAL, "bold": True, "space_after": 2},
        {"t": "730 papers \u00b7 11,130 directed pairs \u00b7 18 Louvain topic clusters", "size": 13, "color": TEXT, "bold": True, "spacing": 1.0},
    ])
    card(s, 7.55, 4.70, 5.2, 1.02, fill=CARD)
    rich(s, 7.75, 4.81, 4.85, 0.85, [
        {"t": "THREE HOMES RECONCILED", "size": 11, "color": TEAL, "bold": True, "space_after": 2},
        {"t": "Zotero 538 \u00b7 group library 911 (27 collections) \u00b7 Airtable 99 \u2192 one audited corpus", "size": 12.5, "color": TEXT, "spacing": 1.0},
    ])
    card(s, 7.55, 5.85, 5.2, 1.02, fill=CARD)
    rich(s, 7.75, 5.96, 4.85, 0.85, [
        {"t": "WORKING BASE", "size": 11, "color": TEAL, "bold": True, "space_after": 2},
        {"t": "957 records: 561 papers \u00b7 39 pathway \u00b7 244 model \u00b7 109 review \u00b7 4 recorder \u2014 43 under the 1,000-record cap", "size": 12, "color": TEXT, "spacing": 1.0},
    ])
    pageno(s, n)

def s27_sadb_curation(prs, n):
    s = content_slide(prs, "Ch. 5 \u00b7 Sensory Afferent Database",
        "The machine drafts at scale; the vocabulary, every judgment call, and every deletion stay human",
        cite="Dissertation Ch. 5 (\u00a7A Structured Afferent Resource, \u00a7AI-Assisted Literature Triage); Bolen et al., SfN 2024",
        notes="Evidence detail: 374 papers machine-drafted in fifteen batches, each batch passing automated vocabulary, "
              "abstract-paste, and duplicate gates before anything was written, and every record verified against its grounding "
              "text after the write (Ch. 5, line 18). Grounding stayed the binding constraint: a PDF-first rule recovered 47 of 55 "
              "flagged papers, an authenticated browser session recovered abstracts for 24 of 70, and the rest stayed bare by rule "
              "(line 18). The 2024 comparative study found AI tools accelerate triage but still require substantial human oversight; "
              "the 2026 campaign is that finding made operational (lines 39-41). Anticipated question: what did the AI actually do? "
              "Answer: batch drafting from grounded text plus the gates; the human held the vocabulary, the 164 flagged judgment "
              "calls, and every single deletion, each prune stage waiting for explicit approval (lines 20, 41).")
    stages = [
        ("RECONCILE", "Zotero 538 + group 911 + Airtable 99, matched record by record on DOI and title"),
        ("IMPORT", "gap imported in verified batches \u00b7 372 auto-classified \u00b7 Papers table 532 \u2192 corpus 943"),
        ("CURATE", "374 papers drafted in 15 gated batches \u00b7 164 flagged judgment calls surfaced for ruling"),
        ("PRUNE", "3 approved stages removed 436 records (382 papers + 54 orphans) \u00b7 351 protected \u00b7 nothing destroyed"),
    ]
    y = 1.90
    for t, sub in stages:
        card(s, 0.55, y, 5.9, 1.02, fill=CARD)
        rich(s, 0.80, y + 0.11, 5.45, 0.85, [
            {"t": t, "size": 11.5, "color": TEAL, "bold": True, "space_after": 2},
            {"t": sub, "size": 12.5, "color": TEXT, "spacing": 1.02},
        ])
        if y < 5.0:
            arrow(s, 3.5, y + 1.04, 3.5, y + 1.22, color=FAINT)
        y += 1.24
    bar_chart(s, 6.85, 1.90, 6.0, 3.35,
              ["notes", "animal-\nprep", "transla-\ntion", "animal-\nstudy", "pathway-\nlinked", "afferent\nclass"],
              [523, 426, 258, 240, 118, 86],
              [GREEN, GREEN, GREEN, GREEN2, GREEN2, GREEN2],
              title="coverage of the 561 retained papers (records carrying each field)",
              num_fmt="0", label_size=12, cat_size=11, max_scale=561)
    card(s, 6.85, 5.45, 6.0, 1.42, fill=CARD2)
    rich(s, 7.05, 5.58, 5.6, 1.2, [
        {"t": "GROUNDING IS THE BINDING CONSTRAINT", "size": 11, "color": TEAL, "bold": True, "space_after": 3},
        {"t": "PDF-first recovery: 47 of 55 flagged papers \u00b7 authenticated browser: abstracts for 24 of 70 \u00b7 the rest stay bare by rule, never guessed",
         "size": 12.5, "color": TEXT, "spacing": 1.05},
    ])
    pageno(s, n)

def s28_sadb_use(prs, n):
    s = content_slide(prs, "Ch. 5 \u00b7 Sensory Afferent Database",
        "A pathway record maps to a circuit edge: class fixes the signal, phase fixes the gating, action fixes the sign",
        cite="Dissertation Ch. 5, Tables 5.1-5.2 (\u00a7Toward Robot and Simulation Use); Ch. 6 implements the mammalian subset",
        notes="Evidence detail: the Feedback table's naming grammar (afferent, action, side, joint, phase) is the same grammar a "
              "circuit edge needs, so translating a record into a controller connection is a lookup rather than a literature search "
              "(Ch. 5, line 45). Linked-paper counts are unchanged between the full audited corpus of 943 and the retained 561 "
              "(Table 5.2 caption, line 64). The 39 records reach beyond the mammalian core - chordotonal organs, trochanteral hair "
              "plates, campaniform sensilla, group III/IV fatigue, and one paper each for somatosensory, vestibular, and visual "
              "postural contributions (line 83). Anticipated question: did the database actually wire the Ch. 6 controller? "
              "Answer: the afferent classes monitored, the signal each channel carries, and the connection rules all trace to "
              "records held here; the controller implements the mammalian subset (lines 12, 45, 83).")
    tx(s, 0.55, 1.86, 5.0, 0.35, "THE CONTROLLER'S SENSORY CHANNELS", size=12, color=TEAL, bold=True)
    chans = [
        ("Ia", "muscle length and velocity", "stretch reflex; phase-dependent modulation"),
        ("II", "static muscle length", "position sense; interjoint coordination"),
        ("Ib", "whole-muscle force", "load-dependent stance regulation"),
        ("heel / toe", "contact events", "discrete stance-swing boundaries"),
    ]
    y = 2.30
    for k, sig, eff in chans:
        card(s, 0.55, y, 5.0, 0.92, fill=CARD)
        tx(s, 0.75, y + 0.09, 1.15, 0.4, k, size=15, color=GREEN, bold=True)
        rich(s, 1.95, y + 0.10, 3.45, 0.78, [
            {"t": sig, "size": 12.5, "color": TEXT, "bold": True, "space_after": 1},
            {"t": eff, "size": 11.5, "color": MUTED, "spacing": 1.0},
        ])
        y += 1.04
    tx(s, 6.05, 1.86, 6.75, 0.35, "PRINCIPAL PATHWAY RECORDS \u00b7 LINKED CORPUS PAPERS", size=12, color=TEAL, bold=True)
    paths = [
        ("Ib stance to swing (the load gate)", 17),
        ("Cutaneous flexor excitation", 16),
        ("Ia or II stance to swing", 10),
        ("Ia monosynaptic excitation", 7),
        ("Ia reciprocal inhibition", 6),
        ("Ib disynaptic excitation (reversal)", 6),
        ("Ib disynaptic inhibition", 5),
        ("Type II excitatory", 5),
    ]
    y = 2.24
    for name, cnt in paths:
        card(s, 6.05, y, 6.75, 0.44, fill=CARD if cnt >= 10 else WHITE)
        tx(s, 6.25, y + 0.07, 5.3, 0.32, name, size=12.5, color=TEXT)
        tx(s, 11.6, y + 0.04, 1.0, 0.32, str(cnt), size=14, color=BROWN, bold=True, align=PP_ALIGN.RIGHT)
        y += 0.52
    lineage = [("Rybak 2006", "two-level CPG"), ("Geyer 2010", "muscle-reflex walking"), ("Deng 2019", "two-layer walker")]
    x = 0.55
    for i, (a, b) in enumerate(lineage):
        node(s, x, 6.44, 2.15, 0.54, a + "\n" + b, fill=CARD2, size=10.5)
        if i < 2:
            arrow(s, x + 2.17, 6.71, x + 2.47, 6.71, color=FAINT, width=1.6)
        x += 2.5
    tx(s, 8.1, 6.46, 4.7, 0.55, "the Models-table lineage the controller instantiates \u00b7 39 pathway records \u00b7 258 robot-and-simulation translation notes",
       size=11, color=MUTED, spacing=1.05)
    pageno(s, n)

# ---------------------------------------------------------------- Ch. 6

def s29_prior_walker(prs, n):
    s = content_slide(prs, "Ch. 6 \u00b7 synthetic nervous system \u00b7 prior work",
        "The direct ancestor: a hand-tuned AnimatLab bipedal walker whose limits set the verification-first agenda",
        cite="Dissertation Ch. 6 \u00a7Prior Work (with C. Morrow); Li et al. 2017; Deng et al. 2019; 2021 NIH BRAIN Initiative Meeting",
        notes="Evidence detail: nine rigid bodies, sagittal-plane hinges, antagonistic Hill-type pairs at the hip, knee, and ankle, "
              "and the author's passive-spring toe; the two-layer CPG plus a sensorimotor layer with Renshaw-like recurrent "
              "inhibition received Ia and Ib feedback, but the gains were never systematically tuned, no ground-contact signals were "
              "modeled, and every current, weight, and tension parameter was set by hand, one at a time (Ch. 6, lines 18-20). "
              "Suspended alternating stepping ran at approximately 4 Hz against roughly 1 Hz for human gait, and was presented at "
              "the 2021 NIH BRAIN Initiative Investigators Meeting (Ch. 7, line 26). Lowered onto a ground plane, the same patterns "
              "could not support the body or generate forward motion (Ch. 7, line 28). The current instrumented model records a "
              "10 s phase-1 run - 50,001 samples at a uniform 0.2 ms interval - demonstrating acquisition, not validated walking "
              "(Ch. 6, line 135). Anticipated question: why keep AnimatLab around? Answer: the lineage's reference models are "
              "archived there and re-run headless by the verification harness, and new circuit organizations are prototyped in the "
              "same format (Ch. 6, line 28).")
    pic(s, os.path.join(FIG, "c_animatlab_body_gui.png"), box=(0.55, 1.86, 6.0, 2.20))
    tx(s, 0.55, 4.12, 6.0, 0.3, "the working project: biomechanical body at initialization", size=11, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    pic(s, os.path.join(FIG, "c_animatlab_network_gui.png"), box=(0.55, 4.50, 6.0, 2.20))
    tx(s, 0.55, 6.76, 6.0, 0.3, "the expanded left-side controller subsystem", size=11, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    rows = [
        ("THE MODEL", "9 rigid bodies \u00b7 sagittal hinges \u00b7 antagonistic Hill-type pairs at hip, knee, ankle \u00b7 passive-spring toe", CARD),
        ("THE CONTROLLER", "two-layer CPG (RG / PF) + sensorimotor layer with Renshaw-like recurrent inhibition \u00b7 Ia and Ib wired", CARD),
        ("THE LIMIT (RED)", "gains never systematically tuned \u00b7 no ground-contact signals \u00b7 every current and weight set by hand", CARD),
    ]
    y = 1.86
    for t, sub, fill in rows:
        card(s, 6.95, y, 5.85, 1.02, fill=fill)
        rich(s, 7.2, y + 0.11, 5.4, 0.85, [
            {"t": t.replace(" (RED)", ""), "size": 11, "color": RED if "RED" in t else TEAL, "bold": True, "space_after": 2},
            {"t": sub, "size": 12.5, "color": TEXT, "spacing": 1.02},
        ])
        y += 1.16
    card(s, 6.95, 5.34, 5.85, 0.80, fill=CARD2)
    tx(s, 7.2, 5.44, 5.4, 0.6, "suspended stepping \u2248 4 Hz (human gait \u2248 1 Hz) \u2014 on the ground it could not support the body",
       size=12.5, color=BROWN, bold=True, spacing=1.02)
    tx(s, 6.95, 6.30, 5.85, 0.65, "hand-tuning consumed the bulk of the effort and did not generalize \u2014 hence verification first",
       size=12, color=MUTED, italic=True, spacing=1.05)
    pageno(s, n)

def s30_toolchain(prs, n):
    s = content_slide(prs, "Ch. 6 \u00b7 synthetic nervous system \u00b7 toolchain",
        "A toolchain that did not exist: audited model conversion, BPA force inside the loop, and SNS blocks in Simulink",
        cite="Dissertation Ch. 6 \u00a7A Neuromechanical Simulation Toolchain; MyoConverter (Caggiano 2022); SNS-Toolbox (Nourse 2023)",
        notes="Evidence detail: no physics engine carried the robot's custom muscle routing together with the BPA force model and an "
              "SNS, so the pipeline is itself a series of tools (Ch. 6, line 12). MyoConverter performs the base conversion; repair "
              "scripts correct mirroring and sign defects and re-audit against the OpenSim source, in a recorded, re-runnable "
              "environment (line 71). All 92 muscle-tendon actuators carried, attachment geometry audited, joint signs verified "
              "dynamically - each actuator activated alone, the sign of its torque identifying its action (line 73). The Simulink "
              "library runs in base Simulink without toolboxes (line 91). Anticipated question: why two engines? Answer: MuJoCo is "
              "the physical engine with the custom routing; Simulink assembles the same component models in a standard "
              "controls-engineering environment - and the rhythm core had to reproduce itself across both before any result was "
              "trusted (lines 91, 113).")
    pic(s, os.path.join(FIG, "c_mujoco_converted_robot.png"), box=(0.55, 1.86, 3.55, 2.85))
    pic(s, os.path.join(FIG, "c_mujoco_synthetic_knee.png"), box=(4.35, 1.86, 3.55, 2.85))
    pic(s, os.path.join(FIG, "c_simulink_sns_library_grid.png"), box=(8.15, 1.86, 4.62, 2.85))
    tx(s, 0.55, 4.78, 3.55, 0.3, "converted robot, reference pose", size=10.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    tx(s, 4.35, 4.78, 3.55, 0.3, "synthetic antagonist-BPA knee, \u221235\u00b0", size=10.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    tx(s, 8.15, 4.78, 4.62, 0.3, "the library's seven core masked blocks", size=10.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    cols = [
        (0.55, 3.55, "OpenSim \u2192 MuJoCo", "92 muscle-tendon actuators \u00b7 attachment geometry audited \u00b7 joint signs verified dynamically"),
        (4.35, 3.55, "BPA force in the loop", "Ch. 3 force surface + Ch. 4 \u03c7 force balance computed in Python at every step"),
        (8.15, 4.62, "Simulink SNS library", "7 core blocks, 12 before the spiking additions \u00b7 knee-reflex demo: sag 2.00\u00b0 with the reflex vs 7.85\u00b0 without"),
    ]
    for x, w, t, sub in cols:
        rich(s, x, 5.16, w, 1.5, [
            {"t": t, "size": 13.5, "color": BROWN, "bold": True, "space_after": 3},
            {"t": sub, "size": 11.5, "color": TEXT, "spacing": 1.05},
        ])
    card(s, 0.55, 6.42, 12.22, 0.62, fill=CARD2)
    tx(s, 0.8, 6.52, 11.8, 0.45,
       "on this coupling: a 410-neuron two-level spinal network \u00b7 2 ms neural step matched to the physics engine \u00b7 isolated rhythm generator oscillates at \u22480.9 s, legs in antiphase",
       size=12, color=TEXT, align=PP_ALIGN.CENTER)
    pageno(s, n)

def s31_rhythm_core(prs, n):
    s = content_slide(prs, "Ch. 6 \u00b7 synthetic nervous system \u00b7 verification",
        "The contact-driven rhythm core, rebuilt in two engines: exact period agreement, trace r \u2265 0.99996",
        cite="Dissertation Ch. 6 \u00a7Cross-Platform Verification (NumPy reference vs native Simulink; 12 s runs at a 2 ms step, 1 Hz heel trains)",
        notes="Evidence detail: the harness is defined around the W2L contact-driven core - four rhythm-generator half-centers built "
              "from spiking neurons whose persistent sodium current terminates each burst, plus graded pattern-formation and sensory "
              "cells - implemented twice from one connectivity specification and driven with an identical protocol (Ch. 6, line 113). "
              "The residual minus 2 ms lag is consistent with the one-step delay of the discrete Simulink solver (line 146). The "
              "headless reruns extract eleven targetable metrics automatically; all eleven pass (line 159); in the modern W2L model "
              "the flexor half-centers peak at minus 56.4 mV and never cross the minus 55 mV threshold, so the rhythm rides graded "
              "extensor-flexor antiphase rather than flexor spiking. Anticipated question: why does cross-engine agreement matter? "
              "Answer: an SNS is a mathematical object that must survive translation between the tools it is built, verified, and "
              "deployed in - the formalism carries the dynamics, not any particular simulator (Ch. 7, line 14).")
    pic(s, os.path.join(FIG, "c_sns_w2l_traces.png"), box=(0.55, 1.86, 5.65, 4.30))
    tx(s, 0.55, 6.22, 5.65, 0.3, "left / right RG extensor voltages: NumPy reference vs Simulink", size=10.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    card(s, 6.45, 1.86, 6.35, 1.98, fill=CARD)
    rich(s, 6.7, 1.98, 5.9, 1.78, [
        {"t": "TWO ENGINES, ONE RHYTHM", "size": 11, "color": TEAL, "bold": True, "space_after": 3},
        {"t": "burst period 1.000 s vs 1.000 s \u2014 0.0% error \u00b7 antiphase r = \u22120.525, exact \u00b7 10 of 10 bursts", "size": 12.5, "color": TEXT, "bold": True, "spacing": 1.05, "space_after": 4},
        {"t": "full traces: r = 0.99996\u20130.99998 at a \u22122 ms lag \u00b7 RMSE 0.099\u20130.111 mV", "size": 12.5, "color": TEXT, "spacing": 1.05},
    ])
    card(s, 6.45, 3.98, 6.35, 2.18, fill=CARD)
    rich(s, 6.7, 4.10, 5.9, 1.98, [
        {"t": "ARCHIVED MODELS RE-RUN HEADLESS \u2014 11 / 11 METRICS PASS", "size": 11, "color": TEAL, "bold": True, "space_after": 3},
        {"t": "bilateral contact-driven: 0.450 s / 0.467 s, lag 0.529 cycle (r = 0.976)", "size": 12, "color": TEXT, "spacing": 1.03, "space_after": 3},
        {"t": "modern W2L: 2.25 Hz \u00b7 flexor peaks \u221256.4 mV never cross \u221255 mV \u00b7 graded antiphase r = \u22120.996", "size": 12, "color": TEXT, "spacing": 1.03, "space_after": 3},
        {"t": "Li walker: 0.637 m/s \u00b7 steps 1.299\u00b10.024 / 1.333\u00b10.080 s \u00b7 duty 0.443 / 0.555", "size": 12, "color": TEXT, "spacing": 1.03},
    ])
    card(s, 0.55, 6.48, 12.22, 0.50, fill=CARD2)
    tx(s, 0.8, 6.56, 11.8, 0.36,
       "first deletion finding: heel trains removed \u2192 static latch in every implementation \u2014 phasic contact feedback is part of the drive architecture, not an optional perturbation",
       size=12, color=BROWN, bold=True, align=PP_ALIGN.CENTER)
    pageno(s, n)

def s32_bpa_walker(prs, n):
    s = content_slide(prs, "Ch. 6 \u00b7 synthetic nervous system \u00b7 end-to-end chain",
        "The characterized BPA model drives a stepping, multi-joint walker in MuJoCo",
        cite="Dissertation Ch. 6 \u00a7BPA-Actuated Walker in MuJoCo \u2014 air-stepping gate; ground failures reported with equal weight",
        notes="Evidence detail: the walker geometry of the prior AnimatLab model exported to MuJoCo and driven by antagonistic BPA "
              "pairs under the split contact-driven network, with actuator forces computed by the Chapter 3 force model (Ch. 6, "
              "line 163). The harness carries 70% of the 411 N body weight; peak BPA forces of 880 N and 1372 N arrive at 372 kPa, "
              "60% of the 620 kPa ceiling, so actuator sizing operates inside the characterized pressure range. Ground negatives: "
              "the W2L gate model stands in full support but does not walk - heel duty zero, forward displacement plus 0.001 m - and "
              "the MuJoCo Li re-implementation produced no stance episodes in any of seven drop-protocol variants (line 165). "
              "Together with the latch result these localize the open problem in the balance and weight-bearing layers, not the "
              "rhythm, the actuators, or the engines (lines 148, 165). Anticipated question: is this walking? Answer: no - it is "
              "marching in a harness that unloads 70% of the body; the chapter reports the ground failures with the same weight as "
              "the air gate. Figure provenance: both knee panels are regenerated from the run artifacts with the column order fixed "
              "by the campaign code (w2l_mujoco/_curriculum_bpa.py JNTS = hip_L, knee_L, ankle_L, hip_R, knee_R, ankle_R), with "
              "knee ranges asserted against each run's own metrics.")
    pic(s, os.path.join(FIG, "c_bpa_ns_knees_big.png"), box=(0.55, 1.88, 2.62, 2.62))
    tx(s, 0.55, 4.40, 2.62, 0.62, "non-spiking \u2014 the chapter's air gate:\nknee flexion + rhythm-generator drive\nartifact: bpa_ns_best_20261003.npz",
       size=9.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER, spacing=1.0)
    pic(s, os.path.join(FIG, "c_bpa_spk_knees_big.png"), box=(3.35, 1.88, 2.62, 2.62))
    tx(s, 3.35, 4.40, 2.62, 0.62, "spiking twin \u2014 since submission:\nknees + rhythm-generator drive\nartifact: bpa_spk_best_20261003.npz",
       size=9.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER, spacing=1.0)
    stats = [
        ("1.542 s", "knee-flexion period"),
        ("63.5\u00b0", "knee-flexion range"),
        ("880 / 1372 N", "peak BPA force, L / R"),
        ("372 kPa", "peak pressure \u2014 60% of the 620 kPa ceiling"),
        ("70%", "of the 411 N body weight, carried by the harness"),
        ("7 / 5", "heel-contact episodes, L / R"),
    ]
    for i, (v, l) in enumerate(stats):
        stat(s, 0.55 + (i % 3) * 1.87, 5.06 + (i // 3) * 0.84, 1.79, v, l, vsize=15, lsize=9)
    card(s, 6.40, 1.88, 6.35, 2.02, fill=CARD)
    rich(s, 6.65, 1.99, 5.9, 1.82, [
        {"t": "ON THE GROUND, SAME WEIGHT", "size": 11, "color": RED, "bold": True, "space_after": 3},
        {"t": "W2L gate model stands (full support) but does not walk: heel duty zero, +0.001 m forward", "size": 12.5, "color": TEXT, "spacing": 1.05, "space_after": 4},
        {"t": "Li re-implementation: no stance episodes in any of seven drop-protocol variants", "size": 12.5, "color": TEXT, "spacing": 1.05},
    ])
    card(s, 6.40, 4.02, 6.35, 1.08, fill=CARD2)
    tx(s, 6.65, 4.13, 5.9, 0.9, "the missing piece is the balance and weight-bearing layer \u2014 not the rhythm, the actuators, or the engines",
       size=13, color=BROWN, bold=True, spacing=1.05)
    card(s, 6.40, 5.22, 6.35, 1.72, fill=CARD, line=TEAL)
    rich(s, 6.65, 5.31, 5.9, 1.58, [
        {"t": "SINCE SUBMISSION", "size": 10.5, "color": TEAL, "bold": True, "space_after": 2},
        {"t": "a spiking twin of this walker now matches the air gate (knee 62.7\u00b0 / 61.9\u00b0, 9 / 9 contact episodes) and outscores the non-spiking build in curriculum tuning (266.9 vs 185.9 over 68 vs 66 trials)",
         "size": 11, "color": TEXT, "spacing": 1.02, "space_after": 2},
        {"t": "artifacts: bpa_spk_best_20261003.npz \u00b7 CURRICULUM_RUN_20261003.log",
         "size": 9, "color": FAINT, "italic": True},
    ])
    pageno(s, n)

def s33_gait_sweep(prs, n):
    s = content_slide(prs, "Ch. 6 \u00b7 synthetic nervous system \u00b7 honest scope check",
        "Tuning to one reference transfers poorly: 43 diverse human reference cycles score one to two orders worse",
        cite="Dissertation Ch. 6 \u00a7Gait-Validation Sweep, Table 6.1 \u2014 library: 32 running + 10 walking + 1 predicted cycle",
        notes="Evidence detail: the library holds 32 running trials from eight subjects at four speeds, ten walking references "
              "(seven prescribed, 0.5-2.0 m/s, plus three self-selected), and Case_40_motion, the Falisse predicted-walking cycle "
              "(period 1.11 s, duty 0.57). The score is the production pipeline's cycle-shape kinematic score: 0 is perfect "
              "agreement and minus 25 is the fixed no-cycle fallback (Ch. 6, line 169). The gap is not an artifact of the running "
              "trials - the production configuration's walking-only mean of minus 8816 is slightly more negative than its "
              "running-only mean of minus 8761 - and all four configurations produce right-leg-only or asymmetric gait (line 187). "
              "For three of four configurations the most negative single reference is Case_40_motion; the W2L-layout variant's is "
              "subject08_Run_20002 (Table 6.1). Anticipated question: why present a mostly negative result? Answer: it quantifies, "
              "honestly, how far single-reference tuning transfers; the family-level signal - the robot-lineage configuration "
              "closest, the most human-derived farthest - is consistent across both evaluations, and closing that gap is a tuning "
              "question, not a tool question (line 187).")
    tx(s, 0.55, 1.84, 4.6, 0.3, "CYCLE-SHAPE KINEMATIC SCORE (0 = PERFECT)", size=11, color=TEAL, bold=True)
    rows = [
        ("w2lvar", "\u2212144.5", "\u22125808", GREEN, "closest"),
        ("s3kpruned", "\u2212176.1", "\u22126924", TEXT, ""),
        ("s3k", "\u2212186.4", "\u22128776", TEXT, ""),
        ("syn6", "\u2212200.4", "\u22129961", RED, "farthest"),
    ]
    y = 2.22
    for name, tr, lm, c, tag in rows:
        card(s, 0.55, y, 4.6, 0.86, fill=CARD)
        tx(s, 0.75, y + 0.08, 1.5, 0.35, name, size=14, color=c, bold=True)
        if tag:
            tx(s, 0.75, y + 0.48, 1.6, 0.3, tag, size=10.5, color=c, bold=True)
        rich(s, 2.35, y + 0.12, 2.7, 0.7, [
            {"t": "tuning-ref  " + tr, "size": 12, "color": MUTED, "space_after": 2},
            {"t": "library mean  " + lm, "size": 13.5, "color": TEXT, "bold": True},
        ])
        y += 1.00
    tx(s, 0.55, 6.12, 4.6, 0.45, "43 references: 32 running \u00b7 10 walking \u00b7 1 predicted", size=10.5, color=MUTED, spacing=1.05)
    pic(s, os.path.join(FIG, "c_w2lvar_joints_big.png"), box=(5.35, 1.86, 3.6, 2.86))
    tx(s, 5.35, 4.78, 3.6, 0.3, "W2L-layout variant \u2014 closest to the library", size=10.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    pic(s, os.path.join(FIG, "c_syn6_joints_big.png"), box=(9.15, 1.86, 3.65, 2.90))
    tx(s, 9.15, 4.82, 3.65, 0.3, "six-synergy variant \u2014 farthest", size=10.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    card(s, 9.15, 5.24, 3.65, 1.55, fill=CARD2)
    rich(s, 9.35, 5.36, 3.3, 1.35, [
        {"t": "NOT A RUNNING ARTIFACT", "size": 10.5, "color": TEAL, "bold": True, "space_after": 2},
        {"t": "s3k on the ten walking refs alone: \u22128816 vs \u22128761 running-only \u2014 all four walk right-leg-only or asymmetric",
         "size": 11.5, "color": TEXT, "spacing": 1.04},
    ])
    pic(s, os.path.join(FIG, "c_s3k_mutfix_skeletal.gif"), box=(5.35, 5.12, 3.4, 1.50))
    tx(s, 5.35, 6.66, 3.4, 0.28, "production s3k, ground evaluation (animated in the talk)", size=10, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    tx(s, 0.55, 6.66, 4.6, 0.3, "the robot-lineage configuration agrees most closely", size=11.5, color=BROWN, bold=True)
    pageno(s, n)

def s34_deletions(prs, n):
    s = content_slide(prs, "Ch. 6 \u00b7 synthetic nervous system \u00b7 stimulus and deletion",
        "The laboratory's first lesion and stimulus experiments read like the physiology they imitate",
        cite="Dissertation Ch. 6 \u00a7Stimulus and Deletion Experiments, Table 6.2; cf. deafferentation: Akay 2014; Lam 2002; Selionov 2009",
        notes="Evidence detail: the sweep removes one pathway class at a time and re-evaluates the tuning-reference score; minus 25 "
              "is the fixed no-cycle fallback (Ch. 6, lines 191-193). Read against the reference literature: deleting all afferent "
              "classes from the production configuration degrades gait without abolishing it, matching the deafferentation outcome "
              "of Akay 2014 - and the six-synergy variant is the network that matches that picture most closely, keeping a "
              "bilateral gait with every afferent class removed. The Ia-only freeze is a stronger dependence than the animal "
              "experiments show and marks that network's tuning as built around its length feedback; the equally large Renshaw "
              "degradation indicates the recurrent-inhibition loop carries weight in the tuned network (line 222). Motoneuron-level "
              "deletion phenomena of the two-level models are not yet reproduced - those and the pattern-formation stimulation "
              "series are future verification tests (lines 191, 222). Anticipated question: which deletion surprised you? Answer: "
              "Renshaw removal hurting as much as group-II removal - recurrent inhibition is usually treated as a footnote.")
    tx(s, 0.55, 1.84, 5.1, 0.3, "PATHWAY REMOVAL \u00b7 PRODUCTION CONFIGURATION (s3k)", size=11, color=TEAL, bold=True)
    rows = [
        ("full model", "\u2212186.4", GREEN, True),
        ("Ia afferents", "\u221225 \u00b7 no cycles detected", RED, False),
        ("all afferent classes", "\u2212267.8", TEXT, False),
        ("Renshaw cells", "\u2212267.3", TEXT, False),
        ("group-II afferents", "\u2212255.2", TEXT, False),
        ("Ib \u00b7 contact \u00b7 interleg commissural", "\u2212180.3 \u00b7 \u2212183.9 \u00b7 \u2212174.9", MUTED, False),
    ]
    y = 2.20
    for name, val, c, strong in rows:
        card(s, 0.55, y, 5.1, 0.44, fill=CARD if strong else WHITE)
        tx(s, 0.75, y + 0.07, 2.9, 0.30, name, size=12, color=TEXT, bold=strong)
        tx(s, 3.35, y + 0.05, 2.1, 0.30, val, size=12, color=c, bold=True, align=PP_ALIGN.RIGHT)
        y += 0.53
    card(s, 6.05, 1.84, 6.75, 1.62, fill=CARD)
    rich(s, 6.3, 1.95, 6.3, 1.42, [
        {"t": "THE VARIANTS DIFFER \u2014 AND THAT IS INFORMATIVE", "size": 11, "color": TEAL, "bold": True, "space_after": 3},
        {"t": "syn6, all afferents removed: \u2212190.5, still bilateral \u2014 10 + 10 cycles \u00b7 knee min \u221271.2\u00b0 \u00b7 0.684 s period", "size": 12, "color": TEXT, "spacing": 1.04, "space_after": 3},
        {"t": "w2lvar needs Ia (\u2212273.5) and RG mutual excitation (\u2212279.9); Renshaw removal leaves \u2212142.4", "size": 12, "color": TEXT, "spacing": 1.04},
    ])
    card(s, 6.05, 3.58, 6.75, 1.62, fill=CARD)
    rich(s, 6.3, 3.69, 6.3, 1.42, [
        {"t": "STIMULUS SIDE \u2014 DRIVE STEPS AND PUSHES", "size": 11, "color": TEAL, "bold": True, "space_after": 3},
        {"t": "s3k walks at +0.01 / \u22120.01 / \u22120.02 (\u2212201.9 / \u2212177.0 / \u2212171.2), freezes at +0.02 \u00b7 15 stand-and-push evaluations: never fell, max sway 0.123 m", "size": 12, "color": TEXT, "spacing": 1.04, "space_after": 3},
        {"t": "syn6 and w2lvar fall under every \u00b10.01 / \u00b10.02 perturbation (pelvis tilt 39\u201340\u00b0)", "size": 12, "color": TEXT, "spacing": 1.04},
    ])
    pic(s, os.path.join(FIG, "c_bilateralrg_sns_diagram.png"), box=(0.55, 5.48, 7.0, 1.46))
    tx(s, 7.75, 5.56, 5.05, 1.4,
       "the walker-network lineage these experiments act on: 110 neurons, 216 synapses \u2014 rendered from the model export with the SNS-Toolbox renderer",
       size=11.5, color=MUTED, spacing=1.1)
    pageno(s, n)

def s35_spiking_mirrors(prs, n):
    s = content_slide(prs, "Ch. 6 \u00b7 synthetic nervous system \u00b7 coding assumption",
        "Spiking mirrors of every simulation family, judged by gates \u2014 not by eyeball",
        cite="Dissertation Ch. 6 \u00a7Spiking-Neuron Mirror Results (gated verification campaign; motor neurons stay non-spiking by design)",
        notes="Evidence detail: the non-spiking baselines were re-established first - all three production configurations reproduce "
              "their documented scores exactly, including minus 161.56754173676563 for the default build, and after the mirror's "
              "selector was added the defaults gate of 410 populations, 376 inputs, and 1186 connections reproduces digit for digit "
              "(Ch. 6, line 232). The topology gate passes exactly under three gain configurations (lines 234). Spikes onto the "
              "non-spiking motor neurons are calibrated by mean conductance with residuals at or below 3e-4 mV (line 121). The "
              "untuned ground gait is a single run with no statistics, claimed only as a lead for the tuning campaign (line 246). "
              "Anticipated question: why does the 55.6% slower rhythm not doom spiking? Answer: the campaign traced the period to "
              "two identifiable knobs - the threshold-increment magnitude and the adaptation time constant - so a period-matched "
              "spiking controller is a short tuning campaign away, and the mapped parameters are the correct starting point "
              "(line 355).")
    pic(s, os.path.join(FIG, "c_spiking_mirror_rhythm.png"), box=(0.55, 1.86, 5.55, 3.90))
    tx(s, 0.55, 5.82, 5.55, 0.3, "constant drive, network alone: analog 0.399 s vs spiking envelope 0.621 s (+55.6%)",
       size=10.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    pic(s, os.path.join(FIG, "c_spiking_mirror_ground_gait.png"), box=(6.35, 1.86, 6.42, 3.90))
    tx(s, 6.35, 5.82, 6.42, 0.3, "ground, same winner parameters: 9 / 0 cycles (left frozen) vs 8 / 9 bilateral, untuned",
       size=10.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    card(s, 0.55, 6.14, 12.22, 0.40, fill=CARD)
    tx(s, 0.8, 6.21, 11.8, 0.28,
       "gates: topology identical (1,186 / 7,908 / 7,762 connections) \u00b7 basin 12/12 under 1% jitter (0.89\u00d7 swing) \u00b7 rhythm missed: +55.6% > 20% \u00b7 air \u224810\u00d7 the 0.47 Hz twin",
       size=10.5, color=TEXT, align=PP_ALIGN.CENTER)
    card(s, 0.55, 6.60, 12.22, 0.44, fill=CARD2)
    tx(s, 0.8, 6.68, 11.8, 0.30,
       "one lead, honestly labeled: untuned mirror \u2212236.35 vs \u2212161.57, knee \u221279.8\u00b0 vs \u221264.3\u00b0 \u2014 bilateral where its twin freezes (single run, no statistics)",
       size=10.5, color=BROWN, bold=True, align=PP_ALIGN.CENTER)
    pageno(s, n)

def s36_what_spiking_changes(prs, n):
    s = content_slide(prs, "Ch. 6 \u00b7 synthetic nervous system \u00b7 coding assumption",
        "Spiking changes timing and robustness; it does not change the interface \u2014 and it forbids parameter copies",
        cite="Dissertation Ch. 6 \u00a7What Spiking Changes and Does Not (knee-reflex and beverage twins; AnimatLab conversions)",
        notes="Evidence detail: with only the sensory interneuron layer converted, the knee-reflex transient is identical - both "
              "read 15.79 degrees at 0.05 s - the settle band moves by 0.9 degrees, and the beverage-holding reflex modulation "
              "moves by 0.02 degrees; to the musculoskeletal model a calibrated 40 Hz train and the graded voltage it replaces are "
              "the same command (Ch. 6, lines 277, 353). The AnimatLab family is the clear negative: no spiking copy preserves the "
              "rhythm, and the thresholds-only ablation isolates the synapse conversion as the change that breaks it - per-spike "
              "conductances at the graded amplitude starve or saturate the reciprocal loop, while the 2023 spiking reference works "
              "because it was tuned as a rate code from birth (lines 327, 359). Anticipated question: do the actuator and torque "
              "conclusions of Chapters 3 and 4 depend on the coding choice? Answer: no - the motor-neuron layer stays non-spiking "
              "and the spikes onto it are calibrated by mean conductance, so the continuous activation interface is unchanged by "
              "construction (line 353).")
    quads = [
        ("UNCHANGED \u2014 THE INTERFACE", GREEN,
         "motor neurons stay graded \u00b7 knee transient identical (15.79\u00b0 at 0.05 s) \u00b7 settle +0.9\u00b0 \u00b7 pour-reflex modulation 5.85\u00b0 vs 5.83\u00b0 (sag 2.00\u00b0 vs 2.02\u00b0 with reflex, 7.85\u00b0 both without)"),
        ("CHANGES \u2014 TIMING", BROWN,
         "constant-drive rhythm 55.6% slower \u00b7 air cadence \u224810\u00d7 faster \u2014 opposite directions, so not one scaling error; traced to two knobs (threshold increment, adaptation constant)"),
        ("CHANGES \u2014 ROBUSTNESS, ITS FAVOR", BLUE,
         "the analog network sits near a bifurcation; the spiking basin passes 12 / 12 at 0.89\u00d7 swing \u2014 the safer substrate for finite-precision hardware, at 0.5 ms sub-steps inside the unchanged 2 ms physics"),
        ("NOT TOLERATED \u2014 NAIVE COPIES", RED,
         "AnimatLab conversions: two-layer dead (2 spikes) \u00b7 bilateral latches at 442\u2013450 Hz \u00b7 reorganized copy co-bursts 11.8 Hz (r = 1.000); thresholds-only keeps oscillating (hip \u221215.1\u00b0 to +23.2\u00b0)"),
    ]
    xs = [0.55, 4.55]
    ys = [1.86, 3.62]
    for i, (t, c, sub) in enumerate(quads):
        x = xs[i % 2]; y = ys[i // 2]
        card(s, x, y, 3.85, 1.64, fill=CARD)
        rich(s, x + 0.2, y + 0.10, 3.5, 1.48, [
            {"t": t, "size": 10.5, "color": c, "bold": True, "space_after": 3},
            {"t": sub, "size": 11, "color": TEXT, "spacing": 1.04},
        ])
    pic(s, os.path.join(FIG, "c_simulink_knee_spiking_twin.png"), box=(8.70, 1.86, 4.10, 2.86))
    tx(s, 8.70, 4.78, 4.10, 0.3, "knee-reflex twin at the matched operating point", size=10.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    card(s, 8.70, 5.16, 4.10, 1.72, fill=CARD2)
    rich(s, 8.90, 5.28, 3.75, 1.52, [
        {"t": "UNIT-LEVEL VERIFICATION", "size": 10.5, "color": TEAL, "bold": True, "space_after": 3},
        {"t": "conductance transient matches the exact exponential to 6.2e-07 \u00b5S \u00b7 hybrid circuit: 62 / 62 spikes, \u226411.75 \u00b5s deviation \u00b7 mean ISI within 0.19 \u00b5s of the toolbox backend",
         "size": 11.5, "color": TEXT, "spacing": 1.05},
    ])
    tx(s, 0.55, 5.50, 7.85, 1.4,
       "the transferable method: calibrate first, gate everything \u2014 synaptic strength in a spiking network is a rate-code quantity derived from firing rate and conductance decay, never copied from the graded value",
       size=13, color=BROWN, bold=True, spacing=1.1)
    pageno(s, n)

# ---------------------------------------------------------------- Ch. 7 / 8 close

def s37_chain(prs, n):
    s = content_slide(prs, "Ch. 7 \u00b7 discussion \u00b7 synthesis",
        "Each stage feeds the next: actuator \u2192 joint \u2192 balance hardware \u2192 afferent database \u2192 neuromechanical laboratory",
        cite="Dissertation Ch. 7 \u00a7Synthesis (chapters 3-6 supply the links)",
        notes="Evidence detail: the force model supplies the actuator terms of the torque model; the torque model supplies the "
              "biomechanical model for designing joint controllers; the validated biomimetic knee shows the architecture reproduces "
              "human joint-level force capability; the database supplies the afferent organization the controller's sensory "
              "channels mirror; and the simulation framework joins the converted model to a synthetic nervous system so "
              "closed-loop experiments can run before hardware exists (Ch. 7, lines 10-12). The cross-engine campaign adds the "
              "design constraint that runs through the chain: an SNS must survive translation between the tools in which it is "
              "built, verified, and deployed - portable primitives, tuned in whichever engine is fastest and trusted in whichever "
              "is most physical (line 14). Anticipated question: what is the single contribution? Answer: the actuator-to-joint "
              "design chain itself - characterize the actuator, correct for transmission compliance, validate against human "
              "biomechanics, and only then close the loop with a controller (Ch. 8, line 8).")
    cards_ = [
        ("ACTUATOR", "Ch. 3", "F*(\u03b5*, P*) surface, three diameter-dependent coefficients; corrects the manufacturer tool at short resting lengths", CARD),
        ("JOINT", "Ch. 4", "\u03c7-corrected torque prediction; biomimetic knee meets or exceeds human isometric torque over most of the range of motion", CARD),
        ("BALANCE", "hardware studies", "a BPA-actuated limb holds itself upright under pseudo-vestibular IMU feedback", CARD),
        ("DATABASE", "Ch. 5", "943-paper audited corpus; the afferent organization the controller's channels mirror", CARD),
        ("SNS LAB", "Ch. 6", "cross-engine trace r > 0.9999 \u00b7 BPA walker peaks at 60% of the pressure ceiling", CARD2),
    ]
    x = MX; cw = 2.36; gap = 0.10
    for i, (k, ch, sub, fill) in enumerate(cards_):
        card(s, x, 2.15, cw, 3.30, fill=fill)
        tx(s, x + 0.16, 2.32, cw - 0.32, 0.32, k, size=11.5, color=TEAL, bold=True)
        tx(s, x + 0.16, 2.66, cw - 0.32, 0.3, ch, size=10.5, color=FAINT, italic=True)
        tx(s, x + 0.16, 3.05, cw - 0.32, 2.3, sub, size=11.5, color=TEXT, spacing=1.1)
        if i < 4:
            arrow(s, x + cw + 0.005, 3.80, x + cw + gap - 0.005, 3.80, color=FAINT, width=2.0)
        x += cw + gap
    card(s, MX, 5.75, 12.22, 0.95, fill=CARD2)
    tx(s, 0.8, 5.88, 11.8, 0.7,
       "the missing piece for unsupported walking is balance \u2014 a network layer, not an engine feature; once it exists, the actuation margin is already there",
       size=14, color=BROWN, bold=True, align=PP_ALIGN.CENTER, spacing=1.05)
    pageno(s, n)

def s38_limits(prs, n):
    s = content_slide(prs, "Ch. 7 \u00b7 discussion \u00b7 honest boundaries",
        "What these results do not claim",
        cite="Dissertation Ch. 7 + chapter-level limitations (Chs. 3-6)",
        notes="These are stated by the dissertation itself. Actuator level: the models are static and isometric - hysteresis, "
              "thermodynamics, and airflow are deliberately set aside (a static model suffices for real-time control; dynamics come "
              "from the Elzein state-space model), and maximum contraction is not predictable a priori, so it is measured per "
              "actuator (Ch. 3, lines 318, 397; Ch. 7, line 6). Joint level: the pinned-flexor data admit a family of stiffness "
              "pairs - the adopted pair was fixed by an external requirement, not by the fit alone - and on the biomimetic "
              "extensor test the corrected prediction still sits above the measurements (FVU 4.49 to 0.67, RMSE 3.03 to 1.17 N m) "
              "(Ch. 4, lines 488, 502). Neural level: full-body co-simulation is blocked by non-finite leg-joint states; the "
              "archived walkers do not walk unsupported; tuned networks transfer poorly to diverse reference cycles; and the "
              "motoneuron-level deletion phenomena of the reference models are not yet reproduced (Ch. 6, lines 73, 165, 187, 363). "
              "The database is deeper than it is broad: the afferent classification covers 86 of 561 retained papers (Ch. 5, "
              "line 87). Anticipated question: which boundary worries you most? Answer: the balance layer - everything else is "
              "engineering.")
    items = [
        ("Static, isometric models", "hysteresis and dynamics set aside by design; \u03b5\u2086\u2082\u2080 not predictable a priori \u2014 measured per actuator (Ch. 3)"),
        ("Stiffness pair not unique", "a family fits the pinned-flexor data; the adopted pair was fixed by an external requirement (Ch. 4)"),
        ("Bio-extensor prediction still high", "corrected FVU 4.49 \u2192 0.67, RMSE 3.03 \u2192 1.17 N\u00b7m \u2014 but it remains above measurement (Ch. 4)"),
        ("Full-body co-simulation blocked", "non-finite leg-joint states; foot-contact meshes the primary suspect (Ch. 6)"),
        ("No unsupported ground walking", "archived walkers stand or latch; the balance layer is not yet present (Ch. 6)"),
        ("Tuning transfers poorly; MN deletions not reproduced", "library scores 1\u20132 orders worse than tuning refs; motoneuron-level phenomena outstanding (Ch. 6)"),
    ]
    x = MX; y = 2.05; cw = 6.0
    for i, (t, sub) in enumerate(items):
        cx = x + (i % 2) * (cw + 0.2)
        cy = y + (i // 2) * 1.50
        card(s, cx, cy, cw, 1.34, fill=CARD)
        rich(s, cx + 0.25, cy + 0.13, cw - 0.5, 1.12, [
            {"t": t, "size": 13.5, "color": BROWN, "bold": True, "space_after": 2},
            {"t": sub, "size": 11.5, "color": TEXT, "spacing": 1.04},
        ])
    tx(s, MX, 6.62, 12.22, 0.4, "and the database is deeper than broad: the afferent classification covers 86 of 561 retained papers (Ch. 5)",
       size=12, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
    pageno(s, n)

def s39_future(prs, n):
    s = content_slide(prs, "Ch. 7 \u00b7 discussion \u00b7 future work",
        "The program the tools define: close the loop, extend the sensorium, build the robot, then maneuver",
        cite="Dissertation Ch. 7 \u00a7Future Work (closing the loop in simulation; sensorium, physical realization, whole-body maneuvers)",
        notes="Evidence detail: the nearest work is the closed loop itself - much of the pipeline is already in place and the "
              "remaining blocker is the non-finite leg-joint states, with the foot-contact meshes the primary suspect; once "
              "resolved, staged validation runs open-loop mechanics first, then a minimal sensor set (contact plus one "
              "proprioceptive channel), then the full Ia/Ib/II set, and only then tuning across walking speeds (Ch. 7, lines 20, "
              "24). The sensorium adds vestibular (IMU) feedback, ocular feedback as a low-frequency correction to the vestibular "
              "estimate, and a cerebellar gateway between commands and the CPG (line 45). Hardware: Teensy microcontrollers for "
              "the per-joint loops, a Jetson-class companion for network evaluation, the static models as feedforward lookup, and "
              "the Elzein pulse-modulated pressure dynamics (lines 47, 49). Anticipated question: what is first? Answer: the "
              "closed loop - the pipeline is assembled and the blocker is specific.")
    stages = [
        ("1 \u00b7 CLOSE THE LOOP", "resolve the non-finite leg-joint blocker \u00b7 staged validation: open-loop \u2192 contact + one proprioceptive channel \u2192 full Ia / Ib / II \u00b7 then tune across speeds", CARD),
        ("2 \u00b7 SENSORIUM", "vestibular (IMU) \u00b7 ocular reset of IMU drift \u00b7 cerebellar gateway between commands and the CPG", CARD),
        ("3 \u00b7 PHYSICAL REALIZATION", "treadmill biped \u00b7 Teensy per-joint loops + Jetson-class SNS evaluation \u00b7 Ch. 3-4 models as feedforward lookup + Elzein pressure dynamics", CARD),
        ("4 \u00b7 WHOLE-BODY MANEUVERS", "gait transitions \u00b7 reactive balancing \u00b7 ballistic, precisely timed bursts \u2014 the kickflip aspiration", CARD2),
    ]
    x = MX; cw = 2.95; gap = 0.14
    for i, (t, sub, fill) in enumerate(stages):
        card(s, x, 2.15, cw, 3.15, fill=fill)
        rich(s, x + 0.2, 2.32, cw - 0.4, 2.85, [
            {"t": t, "size": 12, "color": TEAL, "bold": True, "space_after": 5},
            {"t": sub, "size": 12, "color": TEXT, "spacing": 1.12},
        ])
        if i < 3:
            arrow(s, x + cw + 0.01, 3.70, x + cw + gap - 0.01, 3.70, color=FAINT, width=2.0)
        x += cw + gap
    card(s, MX, 5.62, 12.22, 1.05, fill=CARD2)
    rich(s, 0.8, 5.74, 11.8, 0.85, [
        {"t": "a robot whose nervous system \u2014 unlike a living animal's \u2014 can be instrumented, stimulated, and lesioned at will",
         "size": 15, "color": BROWN, "bold": True, "align": PP_ALIGN.CENTER, "space_after": 3},
        {"t": "the intended demonstration: walking with human-like biomechanics under a biologically grounded neural controller",
         "size": 12, "color": MUTED, "align": PP_ALIGN.CENTER},
    ])
    pageno(s, n)

def s40_ack(prs, n):
    s = content_slide(prs, None,
        "Acknowledgments",
        assert_size=34,
        notes="Keep brief; thank by name. Committee members read the dissertation in full during the revision campaign - say so.")
    people = [
        ("Advisor", "M. Anthony Hunt"),
        ("Committee", "thank you for reading, commenting, and being here"),
        ("Collaborators", "M. Elzein \u00b7 L. Pang \u00b7 L. Burgess \u00b7 C. Morrow \u00b7 J. Scharzenberger \u00b7 C. Steele \u00b7 J. Chung"),
        ("Lab & support", "Agile & Adaptive Robotics Laboratory \u00b7 Portland State University \u00b7 NeuroNex community"),
    ]
    y = 2.3
    for k, v in people:
        rich(s, 1.6, y, 10.2, 1.0, [
            {"t": k.upper(), "size": 12, "color": TEAL, "bold": True, "space_after": 2},
            {"t": v, "size": 17, "color": TEXT},
        ])
        y += 1.12
    pic(s, LOGO, box=(4.87, 6.55, 3.6, 0.8))
    pageno(s, n)

def s41_thanks(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
    pic(s, LOGO, box=(4.47, 1.05, 4.4, 1.34))
    tx(s, 0.7, 3.2, 11.93, 1.1, "Thank you \u2014 questions?", size=44, color=BROWN, bold=True, align=PP_ALIGN.CENTER)
    tx(s, 0.7, 4.5, 11.93, 0.5, "Movement and Control of Biomimetic Humanoid Robots", size=16, color=MUTED, align=PP_ALIGN.CENTER)
    tx(s, 0.7, 5.05, 11.93, 0.5, "Ben P. Bolen  \u00b7  bbolen83@gmail.com", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    tx(s, 0.7, 6.6, 11.93, 0.4, "Appendix: model details, protocol, robustness checks", size=12, color=FAINT, italic=True, align=PP_ALIGN.CENTER)
    s.notes_slide.notes_text_frame.text = (
        "Anticipated questions route to the appendix: force-model coefficients and GoF decomposition (Appendices A-C), "
        "bracket geometry and adopted \u03c7 values (Appendix H), campaign protocol and route provenance (Appendix F), "
        "and the spiking calibration and AnimatLab conversion catalogues (Appendix J).")

def build(prs, start=26):
    s26_sadb_corpus(prs, start + 0)
    s27_sadb_curation(prs, start + 1)
    s28_sadb_use(prs, start + 2)
    s29_prior_walker(prs, start + 3)
    s30_toolchain(prs, start + 4)
    s31_rhythm_core(prs, start + 5)
    s32_bpa_walker(prs, start + 6)
    s33_gait_sweep(prs, start + 7)
    s34_deletions(prs, start + 8)
    s35_spiking_mirrors(prs, start + 9)
    s36_what_spiking_changes(prs, start + 10)
    s37_chain(prs, start + 11)
    s38_limits(prs, start + 12)
    s39_future(prs, start + 13)
    s40_ack(prs, start + 14)
    s41_thanks(prs, start + 15)
    return start + 15
